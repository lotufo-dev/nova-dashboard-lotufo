"""FRC -- Fatura de Recebimento de Consorcio gerada pela dashboard (decisoes do Leonardo, 2026-10-09).

Depois que cada consorciada emite a sua NFS-e contra o consorcio, a FRC (1 pagina) junta essas notas e vai ao
orgao com elas. A FRC e' SO' A SOMA das notas: bruto, INSS, IRRF, ISS e liquido vem da leitura de cada PDF
(nada de cota sobre a medicao), e nota sem leitura confiavel bloqueia a geracao.

As notas esperadas sao as saidas do espelho vigente da natureza da medicao (empresa x municipio, todas as
consorciadas); cada uma e' atendida pelo PDF de nota anexado na medicao (tipo_doc 'nfs') cujo CNPJ do
prestador e municipio, lidos do CONTEUDO, casam com ela -- nunca pelo nome do arquivo.

Uma FRC por linha de `medicoes` (medicao, reajuste e projeto separados). Sai padrao e SEM assinatura, em PDF e
em planilha; a assinada (DocuSign, por fora) entra no envio como documento `fatura` -- a gerada NAO conta como
a FRC do envio (fica em `frcs`, nunca em medicao_documentos). Formulario A = Integra BRT e Complexo Leblon,
desenhado sobre a FRC 21 real (14/09/2026); outro consorcio mostra INDISPONIVEL.
"""
import collections
import datetime
import io
import json
import re
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as ImagemXlsx
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app import arquivos, auditoria, carta, db, espelho, leitura_nfse, leitura_pdf, medicoes, tempo

INDISPONIVEL = "Geração da FRC ainda não disponível para este consórcio"
TRIBUTOS = ("inss", "irrf", "iss")


def _br(centavos: int) -> str:
    return espelho._fmt(centavos)


def _data_br(iso: str) -> str:
    return datetime.date.fromisoformat(iso[:10]).strftime("%d/%m/%Y")


def numero_na_frc(leitura: dict) -> str:
    """Numero da nota como a FRC imprime. O Centi (Guaxe) numera AAAA + 7 digitos ('20260000204'); a FRC 21 real
    escreve '2026204'. Os demais leiautes saem como lidos."""
    s = str(leitura["numero"])
    if leitura.get("leiaute") == "centi" and len(s) == 11 and s[:2] == "20":
        return s[:4] + str(int(s[4:]))
    return s


# --- consorcio, cadastro e numeracao -----------------------------------------------------------------

def _contexto(conn, medicao_id: int) -> tuple[dict, dict | None, dict | None]:
    """(medicao, entidade do consorcio ou None, cadastro da FRC ou None)."""
    m = conn.execute("SELECT * FROM medicoes WHERE id=?", (medicao_id,)).fetchone()
    if m is None:
        raise ValueError("Medição não encontrada")
    ent = conn.execute("SELECT e.* FROM contratos c JOIN entidades e ON e.id=c.entidade_id "
                       "WHERE c.id=? AND e.papel='consorcio'", (m["contrato_id"],)).fetchone()
    if ent is None:
        return dict(m), None, None
    cad = conn.execute("SELECT * FROM frc_consorcios WHERE consorcio=?", (ent["chave"],)).fetchone() \
        if ent["chave"] else None
    return dict(m), dict(ent), ({**dict(cad), "dados": json.loads(cad["dados"])} if cad else None)


def _vigente(conn, medicao_id: int) -> dict | None:
    r = conn.execute("SELECT * FROM frcs WHERE medicao_id=? AND substituida_em IS NULL ORDER BY id DESC LIMIT 1",
                     (medicao_id,)).fetchone()
    return dict(r) if r else None


def ultima(conn, consorcio: str) -> int:
    """A ultima FRC registrada do consorcio: a semente (antes do sistema) ou a maior FRC vigente."""
    r = conn.execute("SELECT MAX(numero) FROM frcs WHERE consorcio=? AND substituida_em IS NULL",
                     (consorcio,)).fetchone()[0]
    semente = conn.execute("SELECT ultimo_anterior FROM frc_consorcios WHERE consorcio=?", (consorcio,)).fetchone()
    return max(r or 0, semente[0] if semente else 0)


# --- notas esperadas e notas anexadas ------------------------------------------------------------------

def _leitura(conn, doc: dict) -> dict:
    """A leitura guardada no anexo (migracao 032). Documento anexado antes dela: le o PDF uma vez e guarda."""
    if doc.get("leitura"):
        return json.loads(doc["leitura"])
    try:
        pdf = arquivos.caminho_seguro(doc["arquivo"]).read_bytes()
        lida = leitura_nfse.ler(leitura_pdf.texto_de_pdf_com_limite(pdf))
    except (OSError, ValueError, leitura_pdf.TextoInadequado):
        return {"ok": False, "motivo": "sem-texto", "leiaute": None}
    with db.mutacao(conn):
        conn.execute("UPDATE medicao_documentos SET leitura=? WHERE id=? AND leitura IS NULL",
                     (json.dumps(lida, ensure_ascii=False), doc["id"]))
    return lida


def valores(l: dict) -> tuple[dict | None, str | None]:
    """Bruto, INSS, IRRF, ISS, retencoes e liquido de uma nota lida -- ou o motivo de nao servir a' FRC.
    Tributo nao lido so' vale zero se o total de retencoes impresso fechar sem ele; e o liquido impresso tem
    que ser o bruto menos as retencoes (nota com PIS/COFINS retidos, por exemplo, nao fecha e bloqueia)."""
    bruto, liquido = l.get("valor_centavos"), l.get("liquido_centavos")
    if bruto is None or liquido is None:
        return None, "valor bruto ou líquido não lido na nota"
    partes = {t: l.get(f"{t}_centavos") for t in TRIBUTOS}
    soma = sum(v or 0 for v in partes.values())
    total = l.get("retencao_total_centavos")
    faltando = [t.upper() for t, v in partes.items() if v is None]
    if faltando and total is None:
        return None, f"{', '.join(faltando)} não lido{'s' if len(faltando) > 1 else ''} na nota"
    if total is not None and total != soma:
        return None, "INSS + IRRF + ISS não fecham com o total de retenções da nota"
    if bruto - soma != liquido:
        return None, "bruto menos retenções não fecha com o líquido da nota"
    return {"bruto": bruto, **{t: v or 0 for t, v in partes.items()}, "retencoes": soma, "liquido": liquido}, None


def _mesmo_municipio(a: str | None, b: str | None) -> bool:
    return bool(a and b) and a.strip().lower() == b.strip().lower()


def checklist(conn, m: dict) -> dict:
    """Saidas esperadas do espelho vigente (todas as empresas) com o PDF que atende cada uma.

    Casamento pelo que foi LIDO: empresa (CNPJ do prestador -> entidade) e municipio; senao empresa e valor
    exato. Mais de um PDF para a mesma saida: vale o ultimo anexado. PDF sem leitura (rascunho sem numero,
    leiaute desconhecido) e PDF lido que nao casa com nenhuma saida aparecem a' parte."""
    from app import notas          # import local: notas importa este modulo
    esp = notas._espelho_vigente(conn, m["id"])
    natureza = next(k for k, v in espelho.TIPO_POR_NATUREZA.items() if v == m["tipo"])
    saidas = [s for s in (esp["saidas"] if esp else []) if s["natureza"] == natureza]
    nomes = {r["chave"]: r["nome_curto"] or r["nome"] for r in conn.execute(
        "SELECT chave, nome, nome_curto FROM entidades WHERE chave IS NOT NULL")}
    por_cnpj = {r["cnpj"]: r["chave"] for r in conn.execute("SELECT cnpj, chave FROM entidades WHERE chave IS NOT NULL")}
    lidas, nao_lidas = [], []
    for d in conn.execute("SELECT id, nome, arquivo, leitura FROM medicao_documentos WHERE medicao_id=? "
                          "AND tipo_doc='nfs' AND arquivo IS NOT NULL ORDER BY id DESC", (m["id"],)):
        l = _leitura(conn, dict(d))
        if l.get("ok"):
            lidas.append({"documento_id": d["id"], "nome": d["nome"], "leitura": l,
                          "empresa": por_cnpj.get(l.get("prestador_cnpj"))})
        else:
            nao_lidas.append({"documento_id": d["id"], "nome": d["nome"], "motivo": l.get("motivo")})
    usados: set[int] = set()
    itens = []
    for s in saidas:
        cand = [x for x in lidas if x["documento_id"] not in usados and x["empresa"] == s["emissor"]]
        doc = next((x for x in cand if _mesmo_municipio(x["leitura"].get("municipio"), s["municipio"])), None) \
            or next((x for x in cand if x["leitura"].get("valor_centavos") == s["valor"]), None)
        item = {"chave": s["chave"], "empresa": nomes.get(s["emissor"], s["emissor"].capitalize()),
                "emissor": s["emissor"], "municipio": s["municipio"], "esperado_centavos": s["valor"],
                "situacao": "falta", "documento_id": None, "nome": None, "numero": None, "valores": None,
                "problema": None}
        if doc:
            usados.add(doc["documento_id"])
            v, problema = valores(doc["leitura"])
            item.update(situacao="anexada", documento_id=doc["documento_id"], nome=doc["nome"],
                        numero=numero_na_frc(doc["leitura"]), valores=v, problema=problema,
                        _leitura=doc["leitura"])
        itens.append(item)
    sobras = [{"documento_id": x["documento_id"], "nome": x["nome"], "numero": numero_na_frc(x["leitura"]),
               "empresa": nomes.get(x["empresa"], "empresa fora do cadastro") if x["empresa"] else "empresa fora do cadastro"}
              for x in lidas if x["documento_id"] not in usados]
    return {"itens": itens, "nao_lidas": nao_lidas, "sobras": sobras, "tem_espelho": esp is not None}


def _bloqueio(ck: dict) -> str | None:
    if not ck["itens"]:
        return ("A medição não tem espelho vigente com as notas esperadas." if not ck["tem_espelho"]
                else "O espelho desta medição não tem notas esperadas para esta natureza.")
    faltam = [f"{i['empresa']} · {i['municipio']}" for i in ck["itens"] if i["situacao"] == "falta"]
    if faltam:
        return f"Falta{'m' if len(faltam) > 1 else ''} {len(faltam)} nota{'s' if len(faltam) > 1 else ''}: " \
               + ", ".join(faltam)
    ruins = [f"NF {i['numero']} ({i['empresa']}): {i['problema']}" for i in ck["itens"] if i["problema"]]
    return "Nota sem leitura confiável — " + "; ".join(ruins) if ruins else None


def _linhas(ck: dict, cad: dict) -> list[dict]:
    """Linhas de INFORMACOES COMPLEMENTARES: por consorciada (ordem do cadastro), depois pelo numero da nota."""
    ordem = {c["chave"]: i for i, c in enumerate(cad["dados"]["consorciadas"])}
    dados = {c["chave"]: c for c in cad["dados"]["consorciadas"]}
    out = []
    for i in ck["itens"]:
        c = dados.get(i["emissor"])
        if c is None:
            raise ValueError(f"Consorciada {i['empresa']} sem CNPJ e razão social no cadastro da FRC")
        out.append({"numero": i["numero"], "cnpj": c["cnpj"], "razao_social": c["razao_social"],
                    "documento_id": i["documento_id"], **i["valores"], "_ordem": (ordem[i["emissor"]], int(i["numero"]))})
    out.sort(key=lambda x: x.pop("_ordem"))
    return out


def _periodo(m: dict, ck: dict) -> tuple[str, str]:
    """Periodo da descricao: o que as notas trazem (o mais comum), senao o mes civil da competencia."""
    lidos = [tuple(i["_leitura"]["periodo"]) for i in ck["itens"] if i.get("_leitura") and i["_leitura"].get("periodo")]
    if lidos:
        return collections.Counter(lidos).most_common(1)[0][0]
    return espelho.periodo_do_mes(m["competencia"])


# --- situacao (o que a tela mostra) ----------------------------------------------------------------------

def situacao(conn, medicao_id: int) -> dict:
    m, ent, cad = _contexto(conn, medicao_id)
    if ent is None:
        return {"disponivel": False, "consorcio": None, "motivo": None}
    if cad is None:
        return {"disponivel": False, "consorcio": ent["nome"], "motivo": INDISPONIVEL}
    ck = checklist(conn, m)
    vig = _vigente(conn, medicao_id)
    ult = ultima(conn, cad["consorcio"])
    frc = None
    if vig:
        linhas_vig = json.loads(vig["notas"])
        atuais = sorted((i["documento_id"], json.dumps(i["valores"], sort_keys=True)) for i in ck["itens"]
                        if i["situacao"] == "anexada" and i["valores"])
        antes = sorted((x["documento_id"], json.dumps({k: x[k] for k in ("bruto", *TRIBUTOS, "retencoes", "liquido")},
                                                      sort_keys=True)) for x in linhas_vig)
        frc = {"numero": vig["numero"], "data_geracao": vig["data_geracao"], "gerada_em": vig["gerada_em"],
               "gerada_por": (conn.execute("SELECT nome FROM usuarios WHERE id=?", (vig["gerada_por"],)).fetchone()
                              or [None])[0],
               "bruto_centavos": vig["bruto_centavos"], "liquido_centavos": vig["liquido_centavos"],
               "desatualizada": atuais != antes}
    itens = [{k: v for k, v in i.items() if not k.startswith("_")} for i in ck["itens"]]
    bloqueio = _bloqueio(ck)
    # Rodada 3 da PR #179 (decisao de produto de 15/10/2026): sem espelho vigente com as notas esperadas e sem FRC gerada,
    # a dashboard nao gera a FRC e a feita a' mao nao vale: o anexo e o envio ficam bloqueados (SEM_ESPELHO).
    doc, recusa = (None, SEM_ESPELHO) if vig is None and not itens else \
        _escolher_assinada(_faturas(conn, medicao_id), vig, bool(frc and frc["desatualizada"]))
    return {"disponivel": True, "consorcio": ent["nome"], "motivo": None, "itens": itens,
            "nao_lidas": ck["nao_lidas"], "sobras": ck["sobras"],
            "completa": bool(itens) and all(i["situacao"] == "anexada" for i in itens),
            "bloqueio": bloqueio, "ultima": ult, "sugerido": vig["numero"] if vig else ult + 1, "frc": frc,
            "assinada": doc is not None, "assinada_documento": doc, "assinada_recusa": recusa,
            "assinada_controla": True}


# --- FRC assinada (issue #169) --------------------------------------------------------------------------------
# A assinada (documento `fatura`) guarda a linha de `frcs` vigente quando foi anexada (migracao 038) e so' vale para
# o envio enquanto essa linha for a vigente e nao estiver desatualizada: a assinada de uma FRC substituida (notas
# trocadas, FRC gerada de novo) nao sai ao orgao com os valores velhos.

# Rodada 3 da PR #179 (decisao de produto de 15/10/2026): FRC feita a' mao so' valeria em contrato marcado "sem formulario",
# e essa marca nao existe no cadastro; entao obra de consorcio sem a FRC no cadastro, ou medicao sem espelho vigente,
# bloqueia o anexo da assinada e o envio.
SEM_CADASTRO = ("A FRC deste consórcio não está no cadastro: a FRC feita à mão não vale para o envio, que fica "
                "bloqueado até o formulário da FRC entrar no cadastro")
SEM_ESPELHO = ("A medição não tem espelho vigente com as notas esperadas: sem ele a dashboard não gera a FRC e a FRC "
               "feita à mão não vale para o envio; gere o espelho e depois a FRC na etapa 3 (Nota fiscal)")
CORROMPIDO = ("O PDF da FRC assinada está inválido ou corrompido: baixe de novo do DocuSign o PDF assinado e anexe "
              "aqui")

DESATUALIZADA = "As notas mudaram depois desta FRC: gere de novo na etapa 3 (Nota fiscal) antes de baixar"


def _faturas(conn, medicao_id: int) -> list[dict]:
    return [dict(r) for r in conn.execute(
        "SELECT id, nome, frc_id FROM medicao_documentos WHERE medicao_id=? AND tipo_doc='fatura' "
        "AND arquivo IS NOT NULL ORDER BY id", (medicao_id,))]


def _escolher_assinada(docs: list[dict], vig: dict | None, desatualizada: bool) -> tuple[dict | None, str | None]:
    """(a assinada que vale, ou None; por que a anexada nao vale, ou None quando nenhuma foi anexada)."""
    validas = [d for d in docs if vig and d["frc_id"] == vig["id"]]
    if validas and not desatualizada:
        return {"id": validas[-1]["id"], "nome": validas[-1]["nome"]}, None
    if not docs:
        return None, None
    if vig is None:
        return None, ("A FRC assinada anexada não corresponde a nenhuma FRC gerada: gere na etapa 3 (Nota fiscal), "
                      "assine no DocuSign e anexe aqui")
    if desatualizada:
        return None, (f"As notas mudaram depois da FRC nº {vig['numero']}: gere de novo na etapa 3 (Nota fiscal), "
                      "assine no DocuSign e anexe aqui")
    return None, (f"A FRC assinada anexada não é a da FRC nº {vig['numero']} vigente: baixe a FRC gerada, assine no "
                  "DocuSign e anexe aqui")


def assinada(conn, medicao_id: int) -> dict:
    """A FRC assinada que vai no envio. `controla`: obra de consorcio, so' vale a assinada da FRC vigente (`recusa`
    diz por que a anexada nao vale, ou por que nenhuma vale: sem espelho, sem a FRC no cadastro). Obra propria: a
    ultima anexada, como antes."""
    s = situacao(conn, medicao_id)
    if s["disponivel"]:
        return {"controla": s["assinada_controla"], "documento": s["assinada_documento"],
                "recusa": s["assinada_recusa"], "frc": s["frc"]}
    if s["consorcio"] is not None:                     # consorcio sem a FRC no cadastro (rodada 3 da PR #179)
        return {"controla": True, "documento": None, "recusa": SEM_CADASTRO, "frc": None}
    docs = _faturas(conn, medicao_id)
    return {"controla": False, "documento": {"id": docs[-1]["id"], "nome": docs[-1]["nome"]} if docs else None,
            "recusa": None, "frc": None}


SEM_TEXTO = ("O PDF da FRC assinada não tem texto para conferir com a FRC gerada: anexe o PDF que o DocuSign entrega, "
             "não uma digitalização")


def _recusa_sem_ler(s: dict) -> dict | None:
    """A recusa que nao precisa do PDF (sem a FRC no cadastro, sem espelho, sem FRC gerada, desatualizada); devolve a
    FRC vigente da situacao `s` quando o PDF tem que ser conferido, None em obra propria."""
    if s["consorcio"] is None:
        return None
    if not s["disponivel"]:
        raise ValueError(SEM_CADASTRO)
    if s["frc"] is None:
        raise ValueError(SEM_ESPELHO if not s["itens"] else
                         "Gere a FRC na etapa 3 (Nota fiscal) antes de anexar a assinada")
    if s["frc"]["desatualizada"]:
        raise ValueError(f"As notas mudaram depois da FRC nº {s['frc']['numero']}: gere de novo na etapa 3 (Nota "
                         "fiscal), assine no DocuSign e anexe a nova")
    return s["frc"]


def precisa_ler(conn, medicao_id: int) -> bool:
    """Antes de ler o PDF e de gravar o arquivo: as recusas baratas sobem aqui (sem gastar uma vaga do leitor nem deixar
    arquivo solto); True quando a assinada tem que ser conferida pelo texto (obra de consorcio)."""
    return _recusa_sem_ler(situacao(conn, medicao_id)) is not None


def ler_assinada(pdf: bytes) -> dict:
    """O texto do PDF assinado, lido FORA da transacao (a leitura roda num subprocesso e leva ate' segundos; dentro do
    BEGIN IMMEDIATE seguraria os outros escritores). A falha fica guardada e sobe em `para_assinar`: a de leitura
    (leitor ocupado, tempo estourado) com a frase dela, que manda tentar de novo; PDF corrompido com CORROMPIDO
    (rodada 3 da PR #179: nada de "PDF inválido (PdfReadError): ..." na tela)."""
    try:
        return {"texto": leitura_pdf.texto_de_pdf_com_limite(pdf), "erro": None}
    except leitura_pdf.TextoInadequado as e:
        return {"texto": None, "erro": e}
    except ValueError:
        return {"texto": None, "erro": ValueError(CORROMPIDO)}


def _no_texto(texto: str, esperado: str) -> bool:
    """`esperado` como palavra inteira no texto, aceitando espaco ou quebra de linha DENTRO dele (o DocuSign quebra
    numero e valor na coluna estreita: "10.000.\n000,07"). Palavra inteira: antes e depois nao pode vir letra, digito,
    ponto, virgula, barra ou hifen; a nota 3456 nao passa por estar dentro da inscricao 123456 (rodada 3 da PR #179)."""
    padrao = r"\s*".join(re.escape(ch) for ch in esperado)
    return re.search(rf"(?<![\w.,/-]){padrao}(?![\w.,/-])", texto) is not None


def esperados(vig: dict) -> list[str]:
    """O que a FRC `vig` (linha de `frcs`) imprime e o PDF assinado tem que trazer: numero, bruto, liquido e o numero
    de cada nota."""
    return [f"{vig['numero']:05d}", _br(vig["bruto_centavos"]), _br(vig["liquido_centavos"]),
            *(str(x["numero"]) for x in json.loads(vig["notas"]))]


def no_texto(texto: str, vig: dict) -> bool:
    """A FRC `vig` inteira no texto (cada item de `esperados` como palavra inteira, ver `_no_texto`)."""
    return all(_no_texto(texto, e) for e in esperados(vig))


def para_assinar(conn, medicao_id: int, lida: dict | None) -> int | None:
    """A FRC (frcs.id) que o PDF assinado assina; `lida` e' o `ler_assinada` do PDF (None em obra propria). Chamada
    DENTRO da transacao do anexo (rodada 3 da PR #179): a FRC gerada de novo ou desatualizada entre a conferencia e o
    INSERT nao deixa o anexo ligado a' versao velha. Sem FRC gerada, desatualizada, sem espelho ou sem a FRC no
    cadastro, recusa. Revisor adversarial na PR #179: a FRC gerada de novo pode manter o numero, entao o PDF tem que
    trazer, no TEXTO, o numero, o total bruto, o total liquido e o numero de cada nota da FRC vigente; a assinada da
    versao anterior (outra nota, outro valor) e' recusada. Obra propria: None (entra sem vinculo, como antes)."""
    atual = _recusa_sem_ler(situacao(conn, medicao_id))
    if atual is None:
        return None
    n = atual["numero"]
    vig = _vigente(conn, medicao_id)
    # Falha de leitura sobe com a frase dela; SEM_TEXTO so' quando o PDF foi lido e nao tem texto (revisor adversarial
    # na PR #179, rodada 2).
    if lida is None:
        raise ValueError("Leia o PDF da FRC assinada antes de conferir")       # defesa: chamador sem ler_assinada
    if lida["erro"] is not None:
        raise lida["erro"]
    texto = lida["texto"]
    if not texto.strip():
        raise ValueError(SEM_TEXTO)
    if not no_texto(texto, vig):
        raise ValueError(f"Este PDF não é a FRC nº {n} vigente (número, notas ou valores não batem): baixe a FRC "
                         "gerada, assine no DocuSign e anexe a assinada dela")
    return vig["id"]


# --- geracao ------------------------------------------------------------------------------------------------

def nome_arquivo(conn, medicao_id: int, ext: str) -> str:
    m, ent, _ = _contexto(conn, medicao_id)
    vig = _vigente(conn, medicao_id)
    return f"FRC {vig['numero']} - {ent['nome'].upper()} ({medicoes.titulo(m).upper()}).{ext}"


def gerar(conn, usuario: dict, medicao_id: int, numero: int) -> dict:
    """Gera (ou gera de novo) a FRC da medicao com o numero dado. Bloqueio -> ValueError com o motivo."""
    m, ent, cad = _contexto(conn, medicao_id)
    if ent is None:
        raise ValueError("A FRC é só de obra de consórcio")
    if cad is None:
        raise ValueError(INDISPONIVEL)
    if m["estado"] in ("cancelada", "desistida"):
        raise ValueError(f"Medição {m['estado']} não gera FRC")
    # Issue #110: o logo vem do cadastro, copiado para a pasta de dados; sem ele a FRC sairia sem o papel timbrado.
    if cad["dados"].get("logo") and carta._arquivo_de_dados(cad["dados"]["logo"]) is None:
        raise ValueError("Logo da FRC não encontrado na pasta de dados: aplique o cadastro")
    if not isinstance(numero, int) or numero < 1:
        raise ValueError("Número da FRC tem que ser um inteiro positivo")
    ck = checklist(conn, m)
    bloqueio = _bloqueio(ck)
    if bloqueio:
        raise ValueError(bloqueio)
    usada = conn.execute("SELECT medicao_id FROM frcs WHERE consorcio=? AND numero=? AND substituida_em IS NULL "
                         "AND medicao_id<>?", (cad["consorcio"], numero, medicao_id)).fetchone()
    if usada:
        outra = conn.execute("SELECT * FROM medicoes WHERE id=?", (usada[0],)).fetchone()
        raise ValueError(f"FRC nº {numero} já é a da {medicoes.titulo(outra)}")
    linhas = _linhas(ck, cad)
    tot = {k: sum(x[k] for x in linhas) for k in ("bruto", *TRIBUTOS, "retencoes", "liquido")}
    if tot["bruto"] - tot["inss"] - tot["irrf"] - tot["iss"] != tot["liquido"]:      # defesa: cada nota ja' fecha
        raise ValueError("Soma das notas não fecha: bruto − INSS − IRRF − ISS ≠ líquido")
    inicio, fim = _periodo(m, ck)
    modelo = cad["dados"]["descricao"].get(m["tipo"])
    if not modelo:
        raise ValueError(f"Cadastro da FRC sem descrição para {medicoes.titulo(m).lower()}")
    hoje = tempo.hoje_local().isoformat()
    doc = {"numero": numero, "data": _data_br(hoje), "dados": cad["dados"], "linhas": linhas, "totais": tot,
           "descricao": modelo.format(n=m["numero"], inicio=_data_br(inicio), fim=_data_br(fim)),
           "nome_consorcio": cad["dados"]["prestador"]["razao_social"]}
    pdf = desenhar_pdf(doc)
    xlsx = desenhar_xlsx(doc)
    rel_pdf, sha_pdf = arquivos.gravar_bytes("frc", pdf, "pdf")          # disco ANTES do banco
    rel_xlsx, sha_xlsx = arquivos.gravar_bytes("frc", xlsx, "xlsx")
    with db.mutacao(conn):
        anterior = _vigente(conn, medicao_id)
        if anterior:
            conn.execute("UPDATE frcs SET substituida_em=datetime('now') WHERE id=?", (anterior["id"],))
        cur = conn.execute(
            "INSERT INTO frcs (medicao_id, consorcio, numero, data_geracao, bruto_centavos, inss_centavos, "
            "irrf_centavos, iss_centavos, liquido_centavos, notas, pdf, pdf_sha256, xlsx, xlsx_sha256, gerada_por) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (medicao_id, cad["consorcio"], numero, hoje, tot["bruto"], tot["inss"], tot["irrf"], tot["iss"],
             tot["liquido"], json.dumps(linhas, ensure_ascii=False), rel_pdf, sha_pdf, rel_xlsx, sha_xlsx,
             usuario["id"]))
        auditoria.registrar(conn, usuario["id"], "frc.gerar", "medicao", medicao_id,
                            {"numero": anterior["numero"]} if anterior else None,
                            {"frc_id": cur.lastrowid, "numero": numero, "consorcio": cad["consorcio"],
                             "bruto": tot["bruto"], "inss": tot["inss"], "irrf": tot["irrf"], "iss": tot["iss"],
                             "liquido": tot["liquido"], "notas": [x["numero"] for x in linhas],
                             "pdf_sha256": sha_pdf, "xlsx_sha256": sha_xlsx})
    return {"id": cur.lastrowid, "numero": numero, **{f"{k}_centavos": v for k, v in tot.items()}}


def gerar_se_completa(conn, usuario: dict, medicao_id: int) -> dict | None:
    """A ultima nota esperada que chega gera a FRC sozinha, com o numero sugerido (ultima + 1). Ja' gerada,
    incompleta ou bloqueada: nada (a tela mostra o motivo)."""
    try:
        s = situacao(conn, medicao_id)
        if not s["disponivel"] or not s["completa"] or s["bloqueio"] or s["frc"]:
            return None
        return gerar(conn, usuario, medicao_id, s["sugerido"])
    except ValueError:
        return None


def arquivo(conn, medicao_id: int, ext: str) -> Path | None:
    vig = _vigente(conn, medicao_id)
    if vig is None:
        return None
    p = arquivos.caminho_seguro(vig["pdf" if ext == "pdf" else "xlsx"])
    return p if p.is_file() else None


# --- desenho: PDF (fpdf2) ----------------------------------------------------------------------------------
# Medidas em mm tiradas da FRC 21 real (planilha da controladoria impressa em PDF, 14/09/2026).

X0, X1 = 17.9, 189.8
CINZA = (217, 217, 217)
COL_SERV = (17.9, 27.0, 36.0, 127.5, 136.5, 145.5, 167.7, 189.8)      # SEQ | (vazia) | DESCRICAO | UM | QTDE | VU | VT
COL_INFO = (17.9, 36.0, 63.0, 127.5, 145.5, 167.7, 189.8)             # NFSE | CNPJ | RAZAO | BRUTO | RET | LIQ


class _Pdf(FPDF):
    def __init__(self):
        super().__init__(unit="mm", format="A4")
        self.set_auto_page_break(False)
        self.set_margins(0, 0, 0)
        if carta._ARIAL.is_file() and carta._ARIAL_NEGRITO.is_file():
            self.add_font("Frc", "", str(carta._ARIAL))
            self.add_font("Frc", "B", str(carta._ARIAL_NEGRITO))
            self.fonte, self.unicode = "Frc", True
        else:
            self.fonte, self.unicode = "Helvetica", False
        self.set_line_width(0.2)
        self.set_draw_color(0, 0, 0)

    def t(self, s: str) -> str:
        if self.unicode:
            return s
        for de, para in carta._TROCAS_LATIN1.items():
            s = s.replace(de, para)
        return s.encode("latin-1", "replace").decode("latin-1")

    def texto(self, x: float, y: float, w: float, s: str, tam: float, negrito: bool = False, align: str = "L"):
        """Uma linha com o centro vertical em `y`."""
        self.set_font(self.fonte, "B" if negrito else "", tam)
        self.set_xy(x, y - tam * 0.2)
        self.cell(w, tam * 0.4, self.t(s), align=align)

    def rotulo_valor(self, x: float, y: float, rotulo: str, valor: str, tam: float = 6.1):
        self.set_font(self.fonte, "", tam)
        larg = self.get_string_width(self.t(rotulo + " "))
        self.texto(x, y, larg, rotulo + " ", tam)
        self.texto(x + larg, y, 120, valor, tam, negrito=True)

    def caixa(self, x0: float, y0: float, x1: float, y1: float, fundo: bool = False):
        if fundo:
            self.set_fill_color(*CINZA)
        self.rect(x0, y0, x1 - x0, y1 - y0, style="DF" if fundo else "D")


def _bloco_dados(pdf: _Pdf, y_barra: float, y_corpo: float, y_fim: float, titulo: str, d: dict, linhas: list[float]):
    pdf.caixa(X0, y_barra, X1, y_corpo, fundo=True)
    pdf.texto(X0, (y_barra + y_corpo) / 2, X1 - X0, titulo, 6.3, True, "C")
    pdf.caixa(X0, y_corpo, X1, y_fim)
    esq = [("Nome/Razão Social:", d["razao_social"]), ("ENDEREÇO:", d["endereco"]),
           ("MUNICÍPIO/ESTADO/PAIS:", d["municipio"]), ("CNPJ:", d["cnpj"]), ("TELEFONE:", d["telefone"])]
    for y, (r, v) in zip(linhas, esq):
        pdf.rotulo_valor(18.5, y, r, v)
    dir_ = [("CEP:", d["cep"]), ("INSCRIÇÃO MUNICIPAL:", d["inscricao_municipal"]), ("E-MAIL:", d["email"])]
    for y, (r, v) in zip(linhas[2:], dir_):
        pdf.rotulo_valor(128.0, y, r, v)


def desenhar_pdf(doc: dict) -> bytes:
    d = doc["dados"]
    pdf = _Pdf()
    pdf.add_page()
    # cabecalho: logo | titulo e identificacao | (QR da planilha original: nao reproduzido)
    pdf.caixa(X0, 19.2, X1, 43.2)
    pdf.line(72.1, 19.2, 72.1, 43.2)
    pdf.line(127.5, 19.2, 127.5, 43.2)
    logo = carta._arquivo_de_dados(d.get("logo"))   # copiado do cadastro para a pasta de dados (issue #110)
    if logo and logo.is_file():
        pdf.image(str(logo), x=20.65, y=19.4, w=48.6)
    else:
        pdf.set_font(pdf.fonte, "B", 8)
        pdf.set_xy(X0 + 2, 26)
        pdf.multi_cell(72.1 - X0 - 4, 4, pdf.t(doc["nome_consorcio"]), align="C")
    pdf.texto(72.1, 21.65, 55.4, "FATURA DE RECEBIMENTO DE CONSÓRCIO", 7.2, True, "C")
    pdf.set_fill_color(*CINZA)
    pdf.rect(72.1, 24.6, 55.4, 13.35, style="F")
    colunas = ((74.1, 18.4), (90.8, 18.4), (107.4, 18.4))
    for (x, w), (r1, v1), (r2, v2) in zip(colunas, (("Número da Fatura", f"{doc['numero']:05d}"),
                                                    ("Série", "ÚNICA"), ("Data de Geração", doc["data"])),
                                          (("Natureza da Operação", "RECEBIMENTO"), ("Sigla", "FRC"),
                                           ("Data de Vencimento", "25 DDL"))):
        pdf.texto(x, 25.75, w, r1, 4.6, False, "C")
        pdf.texto(x, 28.4, w, v1, 5.8, True, "C")
        pdf.texto(x, 33.8, w, r2, 4.6, False, "C")
        pdf.texto(x, 36.4, w, v2, 5.8, True, "C")
    pdf.set_font(pdf.fonte, "", 4.6)
    rot = pdf.t("Código Autenticador Interno: ")
    larg = pdf.get_string_width(rot)
    pdf.set_font(pdf.fonte, "B", 5.2)
    larg_v = pdf.get_string_width("GL.SGI.FRP.300")
    x = 72.1 + (55.4 - larg - larg_v) / 2
    pdf.texto(x, 40.5, larg, rot, 4.6)
    pdf.texto(x + larg, 40.5, larg_v + 1, "GL.SGI.FRP.300", 5.2, True)

    _bloco_dados(pdf, 43.2, 48.55, 67.2, "DADOS DO PRESTADOR DE SERVIÇOS", d["prestador"],
                 [52.4, 55.07, 57.74, 60.41, 63.08])
    _bloco_dados(pdf, 67.2, 72.5, 91.2, "DADOS DO TOMADOR DE SERVIÇOS", d["tomador"],
                 [76.5, 79.17, 81.84, 84.51, 87.18])

    # servico
    pdf.caixa(X0, 91.2, X1, 96.5, fundo=True)
    for x in COL_SERV[1:-1]:
        pdf.line(x, 91.2, x, 114.2)
    for (a, b), s in zip(zip(COL_SERV, COL_SERV[1:]), ("SEQ.", "", "DESCRIÇÃO DOS SERVIÇOS", "U.M.", "QTDE.",
                                                         "V. UNITÁRIO", "VALOR TOTAL")):
        pdf.texto(a, 93.85, b - a, s, 6.3, True, "C")
    pdf.caixa(X0, 96.5, X1, 114.2)
    meio = (96.5 + 114.2) / 2
    pdf.texto(COL_SERV[0], meio, COL_SERV[1] - COL_SERV[0], "1", 6.1, False, "C")
    pdf.set_font(pdf.fonte, "", 5.8)
    linhas_desc = pdf.multi_cell(COL_SERV[3] - COL_SERV[2] - 0.4, 2.6, pdf.t(doc["descricao"]), align="C",
                                 dry_run=True, output="LINES")
    pdf.set_xy(COL_SERV[2] + 0.2, meio - len(linhas_desc) * 2.6 / 2)
    pdf.multi_cell(COL_SERV[3] - COL_SERV[2] - 0.4, 2.6, pdf.t(doc["descricao"]), align="C")
    pdf.texto(COL_SERV[3], meio, COL_SERV[4] - COL_SERV[3], "UND.", 6.1, False, "C")
    pdf.texto(COL_SERV[4], meio, COL_SERV[5] - COL_SERV[4], "1", 6.1, False, "C")
    tot = doc["totais"]
    for (a, b) in ((COL_SERV[5], COL_SERV[6]), (COL_SERV[6], COL_SERV[7])):
        pdf.texto(a + 1.6, meio, 8, "R$", 5.8)
        pdf.texto(a, meio, b - a - 1.3, _br(tot["bruto"]), 5.8, False, "R")

    # totais
    y = 114.2
    for rotulo, chave in (("TOTAL BRUTO FATURADO", "bruto"), ("RETENÇÃO DE INSS", "inss"), ("RETENÇÃO DE IRRF", "irrf"),
                          ("RETENÇÃO DE ISS", "iss"), ("TOTAL LIQUIDO FATURADO", "liquido")):
        pdf.caixa(X0, y, X1, y + 5.08, fundo=True)
        pdf.line(167.7, y, 167.7, y + 5.08)
        pdf.texto(X0, y + 2.54, 167.7 - X0 - 0.6, rotulo, 6.3, True, "R")
        pdf.texto(169.3, y + 2.54, 8, "R$", 5.6, True)
        pdf.texto(167.7, y + 2.54, X1 - 167.7 - 1.3, _br(tot[chave]), 5.6, True, "R")
        y += 5.08

    # informacoes complementares
    y = 142.3
    pdf.caixa(X0, y, X1, y + 2.65, fundo=True)
    pdf.texto(X0, y + 1.33, X1 - X0, "INFORMAÇÕES COMPLEMENTARES", 6.3, True, "C")
    y += 2.65
    cab = ("N° NFSE", "CNPJ", d.get("rotulo_empresa") or "RAZÃO SOCIAL", "V. BRUTO", "RETENÇÕES", "V. LIQUIDO")
    linhas = [cab] + [(x["numero"], x["cnpj"], x["razao_social"], _br(x["bruto"]), _br(x["retencoes"]), _br(x["liquido"]))
                      for x in doc["linhas"]]
    for i, linha in enumerate(linhas):
        pdf.caixa(X0, y, X1, y + 3.56)
        for x in COL_INFO[1:-1]:
            pdf.line(x, y, x, y + 3.56)
        for (a, b), s in zip(zip(COL_INFO, COL_INFO[1:]), linha):
            pdf.texto(a, y + 1.78, b - a, str(s), 4.4, i == 0, "C")
        y += 3.56

    # linhas para as assinaturas (a FRC sai sem assinatura: o DocuSign assina por fora)
    pdf.line(72.1, 183.3, 127.5, 183.3)
    pdf.line(X0, 188.6, 63.0, 188.6)
    pdf.line(136.5, 188.6, X1, 188.6)
    return bytes(pdf.output())


# --- desenho: planilha (openpyxl) ---------------------------------------------------------------------------
# A FRC original e' uma planilha: a grade abaixo segue as mesmas colunas do PDF (A..K) e as mesmas alturas.

_LARGURAS_MM = (9.1, 9.0, 27.0, 9.1, 18.5, 18.4, 18.5, 9.0, 9.0, 22.2, 22.1)
_FINO = Side(style="thin", color="000000")
_CINZA = PatternFill("solid", fgColor="D9D9D9")
# A grade e as fontes vao 1,5x maiores que no PDF (Arial no Excel pede folga) e o "ajustar a 1 pagina" devolve ao A4.
ESCALA = 1.5


def desenhar_xlsx(doc: dict) -> bytes:
    d = doc["dados"]
    wb = Workbook()
    ws = wb.active
    ws.title = "FRC"
    ws.sheet_view.showGridLines = False
    for i, mm in enumerate(_LARGURAS_MM):
        ws.column_dimensions[chr(65 + i)].width = round((mm * ESCALA * 96 / 25.4 - 5) / 7, 2)   # unidade do Excel

    def alt(linha: int, mm: float):
        ws.row_dimensions[linha].height = round(mm * ESCALA * 2.835, 1)

    def cel(ref: str, valor, tam: float = 8, negrito: bool = False, h: str = "center", fundo: bool = False,
            quebra: bool = False, fmt: str | None = None):
        c = ws[ref]
        c.value = valor
        c.font = Font(name="Arial", size=tam, bold=negrito)
        c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=quebra)
        if fundo:
            c.fill = _CINZA
        if fmt:
            c.number_format = fmt
        return c

    def juntar(faixa: str):
        ws.merge_cells(faixa)

    def moldura(faixa: str, fundo: bool = False):
        linhas = ws[faixa if ":" in faixa else f"{faixa}:{faixa}"]
        for r, linha in enumerate(linhas):
            for k, c in enumerate(linha):
                c.border = Border(left=_FINO if k == 0 else None, right=_FINO if k == len(linha) - 1 else None,
                                  top=_FINO if r == 0 else None, bottom=_FINO if r == len(linhas) - 1 else None)
                if fundo:
                    c.fill = _CINZA

    def rico(rotulo: str, valor: str, tam: float = 9):
        return CellRichText(TextBlock(InlineFont(rFont="Arial", sz=tam), rotulo + " "),
                            TextBlock(InlineFont(rFont="Arial", sz=tam, b=True), valor))

    # cabecalho (linhas 1-7)
    for linha, mm in zip(range(1, 8), (5.4, 2.6, 3.0, 1.4, 2.6, 3.0, 5.2)):
        alt(linha, mm)
    juntar("A1:D7"); moldura("A1:D7")
    juntar("H1:K7"); moldura("H1:K7")
    moldura("E1:G7")
    juntar("E1:G1"); cel("E1", "FATURA DE RECEBIMENTO DE CONSÓRCIO", 10.5, True)
    for col, (r1, v1), (r2, v2) in zip("EFG", (("Número da Fatura", f"{doc['numero']:05d}"), ("Série", "ÚNICA"),
                                               ("Data de Geração", doc["data"])),
                                       (("Natureza da Operação", "RECEBIMENTO"), ("Sigla", "FRC"),
                                        ("Data de Vencimento", "25 DDL"))):
        cel(f"{col}2", r1, 7, fundo=True); cel(f"{col}3", v1, 8.5, True, fundo=True)
        ws[f"{col}4"].fill = _CINZA
        cel(f"{col}5", r2, 7, fundo=True); cel(f"{col}6", v2, 8.5, True, fundo=True)
    juntar("E7:G7"); ws["E7"].value = rico("Código Autenticador Interno:", "GL.SGI.FRP.300", 7)
    ws["E7"].alignment = Alignment(horizontal="center", vertical="center")
    logo = carta._arquivo_de_dados(d.get("logo"))   # copiado do cadastro para a pasta de dados (issue #110)
    if logo and logo.is_file():
        img = ImagemXlsx(str(logo))
        px = ESCALA * 96 / 25.4                         # pixels por mm na grade
        img.width, img.height = round(44 * px), round(44 * px * img.height / img.width)
        emu = 9525                                       # EMU por pixel
        img.anchor = OneCellAnchor(_from=AnchorMarker(col=0, row=0, colOff=round(5 * px) * emu, rowOff=round(0.6 * px) * emu),
                                   ext=XDRPositiveSize2D(img.width * emu, img.height * emu))
        ws.add_image(img)
    else:
        cel("A1", doc["nome_consorcio"], 11, True, quebra=True)

    def dados(inicio: int, titulo: str, p: dict) -> int:
        alt(inicio, 5.35)
        juntar(f"A{inicio}:K{inicio}"); moldura(f"A{inicio}:K{inicio}", fundo=True)
        cel(f"A{inicio}", titulo, 9.5, True, fundo=True)
        alt(inicio + 1, 1.2)
        esq = [("Nome/Razão Social:", p["razao_social"]), ("ENDEREÇO:", p["endereco"]),
               ("MUNICÍPIO/ESTADO/PAIS:", p["municipio"]), ("CNPJ:", p["cnpj"]), ("TELEFONE:", p["telefone"])]
        dir_ = [None, None, ("CEP:", p["cep"]), ("INSCRIÇÃO MUNICIPAL:", p["inscricao_municipal"]), ("E-MAIL:", p["email"])]
        for k, (e, dd) in enumerate(zip(esq, dir_)):
            r = inicio + 2 + k
            alt(r, 3.3)          # acento de maiuscula (Ó, Á) pede a linha mais alta que no PDF
            ws[f"A{r}"].value = rico(*e)
            if dd:
                ws[f"H{r}"].value = rico(*dd)
        alt(inicio + 7, 1.8)
        moldura(f"A{inicio + 1}:K{inicio + 7}")
        return inicio + 8

    r = dados(8, "DADOS DO PRESTADOR DE SERVIÇOS", d["prestador"])
    r = dados(r, "DADOS DO TOMADOR DE SERVIÇOS", d["tomador"])

    # servico
    alt(r, 5.3)
    for faixa, s in (("A", "SEQ."), ("B", ""), ("C:G", "DESCRIÇÃO DOS SERVIÇOS"), ("H", "U.M."), ("I", "QTDE."),
                     ("J", "V. UNITÁRIO"), ("K", "VALOR TOTAL")):
        a, _, b = faixa.partition(":")
        if b:
            juntar(f"{a}{r}:{b}{r}")
        moldura(f"{a}{r}:{b or a}{r}", fundo=True)
        cel(f"{a}{r}", s, 9.5, True, fundo=True)
    r += 1
    alt(r, 17.7)
    juntar(f"C{r}:G{r}")
    for faixa in ("A", "B", "C:G", "H", "I", "J", "K"):
        a, _, b = faixa.partition(":")
        moldura(f"{a}{r}:{b or a}{r}")
    tot = doc["totais"]
    cel(f"A{r}", 1, 8.5); cel(f"C{r}", doc["descricao"], 8.5, quebra=True); cel(f"H{r}", "UND.", 8.5)
    cel(f"I{r}", 1, 8.5)
    moeda = '"R$" * #,##0.00'
    cel(f"J{r}", tot["bruto"] / 100, 8.5, h="right", fmt=moeda); cel(f"K{r}", tot["bruto"] / 100, 8.5, h="right", fmt=moeda)
    r += 1
    for rotulo, chave in (("TOTAL BRUTO FATURADO", "bruto"), ("RETENÇÃO DE INSS", "inss"), ("RETENÇÃO DE IRRF", "irrf"),
                          ("RETENÇÃO DE ISS", "iss"), ("TOTAL LIQUIDO FATURADO", "liquido")):
        alt(r, 5.08)
        juntar(f"A{r}:J{r}"); moldura(f"A{r}:J{r}", fundo=True); moldura(f"K{r}", fundo=True)
        cel(f"A{r}", rotulo, 9.5, True, h="right", fundo=True)
        cel(f"K{r}", tot[chave] / 100, 8.5, True, h="right", fundo=True, fmt=moeda)
        r += 1
    alt(r, 2.7)
    r += 1

    # informacoes complementares
    alt(r, 3.2)
    juntar(f"A{r}:K{r}"); moldura(f"A{r}:K{r}", fundo=True)
    cel(f"A{r}", "INFORMAÇÕES COMPLEMENTARES", 9.5, True, fundo=True)
    r += 1
    faixas = ("A:B", "C", "D:G", "H:I", "J", "K")
    cab = ("N° NFSE", "CNPJ", d.get("rotulo_empresa") or "RAZÃO SOCIAL", "V. BRUTO", "RETENÇÕES", "V. LIQUIDO")
    linhas = [cab] + [(x["numero"], x["cnpj"], x["razao_social"], x["bruto"] / 100, x["retencoes"] / 100,
                       x["liquido"] / 100) for x in doc["linhas"]]
    for i, linha in enumerate(linhas):
        alt(r, 3.56)
        for faixa, v in zip(faixas, linha):
            a, _, b = faixa.partition(":")
            if b:
                juntar(f"{a}{r}:{b}{r}")
            moldura(f"{a}{r}:{b or a}{r}")
            cel(f"{a}{r}", v, 6.6, i == 0, fmt="#,##0.00" if isinstance(v, float) else None)
        r += 1

    ws.print_area = f"A1:K{r + 6}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
