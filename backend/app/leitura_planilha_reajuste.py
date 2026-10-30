"""Leitura da planilha de previsao de reajuste do pleito (issue #272; spec 2026-10-20 §2.2 e §4 item 1).

O modelo "Mem. e Cron. - Prev. de Reajuste" tem a aba `Resumo` com uma linha por medicao futura: cabecalho
MEDICAO | PERIODO | VALORES e, na linha de baixo, A PI | (INDICE) | REAJUSTE | TOTAL. A coluna inicial varia
entre as obras (A ou B) e o INDICE nem sempre existe, entao as colunas sao achadas pelo texto do cabecalho.
A competencia e' o mes do FIM do periodo ("23/02/2026 a 31/03/2026" e' 03/2026: a 1a medicao pode comecar no
mes anterior). O valor e' a coluna REAJUSTE, em reais, que a planilha ja' trunca em 2 casas.

Deterministico (openpyxl, sem custo de API). O que nao casa com o modelo levanta `PlanilhaNaoReconhecida` e a
pessoa digita a previsao (fila humana, origem `manual`). Nada aqui grava: quem grava e' `pleitos`.
"""
import datetime
import io
import re
import unicodedata
import zipfile

import openpyxl
from defusedxml import ElementTree       # XML de fora: sem entidade expandida (o openpyxl usa o mesmo)
from openpyxl.utils import get_column_letter

from app import dinheiro, rm

ABA = "Resumo"
MENSAGEM = "Não reconheci a planilha; digite a previsão"
LINHAS_CABECALHO = 40          # o cabecalho do modelo fica nas primeiras linhas (6 e 7, ou 3 e 4)
MAX_COMPETENCIAS = 240         # 20 anos de medicoes mensais: acima disso nao e' o modelo (pleitos usa o mesmo)
MAX_VARRIDAS = 5000            # linhas lidas da aba (as em branco formatadas contam aqui, nao no teto acima)
MAX_COLUNAS = 60               # o cabecalho do modelo ocupa 4 a 7 colunas; celula distante nao entra na leitura
# teto do XML descompactado da aba Resumo e de cada parte que o openpyxl le inteira ao abrir (estilos, textos): o
# .xlsx pequeno com milhoes de linhas vazias comprimidas nao pode custar memoria e CPU do servidor. A aba Resumo do
# modelo tem ~40 linhas; formatada ate' a linha 8000 em 20 colunas, fica perto de 4 MB
MAX_BYTES_XML = 10_000_000
_DATA = re.compile(r"([0-9]{2})/([0-9]{2})/([0-9]{4})")


class PlanilhaNaoReconhecida(ValueError):
    """A planilha nao casa com o modelo. `str()` e' a mensagem da tela; `motivo` diz o que faltou."""

    def __init__(self, motivo: str):
        super().__init__(MENSAGEM)
        self.motivo = motivo


def _texto(v) -> str:
    """Cabecalho comparavel: sem acento, maiusculo, sem espacos nas pontas ("MEDIÇÃO " -> "MEDICAO")."""
    if not isinstance(v, str):
        return ""
    s = unicodedata.normalize("NFKD", v).encode("ascii", "ignore").decode()
    return " ".join(s.upper().split())


def _competencia(periodo) -> str | None:
    """Mes do fim do periodo: "01/04/2026 a 30/04/2026" -> "2026-04". Data do Excel tambem vale."""
    if isinstance(periodo, datetime.datetime | datetime.date):
        return f"{periodo.year:04d}-{periodo.month:02d}"
    datas = _DATA.findall(periodo) if isinstance(periodo, str) else []
    if not datas:
        return None
    dia, mes, ano = datas[-1]
    try:
        datetime.date(int(ano), int(mes), int(dia))
    except ValueError:
        return None
    return f"{ano}-{mes}"


def _centavos(v) -> int:
    """Reais da planilha (float) -> centavos inteiros: o arredondamento unico do projeto (`rm.centavos`, pelo texto
    do float: 21342.29 nao vira 2134228)."""
    c = rm.centavos(v)
    if abs(c) > dinheiro.TETO_CENTAVOS:
        raise PlanilhaNaoReconhecida("valor de reajuste acima do máximo aceito")
    return c


def _colunas(ws) -> tuple[int, int, int, int]:
    """(linha do primeiro dado, coluna MEDICAO, coluna PERIODO, coluna REAJUSTE), achadas pelo cabecalho."""
    linhas = list(ws.iter_rows(min_row=1, max_row=LINHAS_CABECALHO, max_col=MAX_COLUNAS, values_only=True))
    for i, linha in enumerate(linhas):
        textos = [_texto(v) for v in linha]
        if "MEDICAO" not in textos or "PERIODO" not in textos:
            continue
        col_med, col_per = textos.index("MEDICAO"), textos.index("PERIODO")
        # REAJUSTE na mesma linha ou na de baixo (o modelo agrupa os valores sob "VALORES")
        for j in (i, i + 1):
            sub = [_texto(v) for v in linhas[j]] if j < len(linhas) else []
            if "REAJUSTE" in sub:
                return j + 2, col_med, col_per, sub.index("REAJUSTE")
        raise PlanilhaNaoReconhecida("a aba Resumo não tem a coluna REAJUSTE")
    raise PlanilhaNaoReconhecida("a aba Resumo não tem o cabeçalho MEDIÇÃO e PERÍODO")


def ler(dados: bytes) -> list[dict]:
    """Previsao por competencia: [{"competencia": "AAAA-MM", "valor_centavos": int, "medicao", "periodo"}]."""
    try:
        _recusar_parte_grande(dados)
        # read_only: so' a aba Resumo e' lida (as abas de memoria do modelo tem milhares de colunas)
        wb = openpyxl.load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
    except zipfile.BadZipFile as e:
        raise PlanilhaNaoReconhecida("o arquivo não abre como planilha .xlsx") from e
    except PlanilhaNaoReconhecida:
        raise
    except Exception as e:   # arquivo de fora: XML quebrado (ParseError e afins) vira digitacao, nao erro 500
        raise PlanilhaNaoReconhecida("a planilha não pôde ser lida") from e
    try:
        return _ler_resumo(wb, dados)
    except PlanilhaNaoReconhecida:
        raise
    except Exception as e:   # arquivo de fora: XML quebrado dentro da aba vira digitacao, nao erro 500
        raise PlanilhaNaoReconhecida("a aba Resumo não pôde ser lida") from e
    finally:
        wb.close()


_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _recusar_parte_grande(dados: bytes) -> None:
    """Antes de abrir: as partes XML que o openpyxl le inteiras (tudo menos as abas, lidas sob demanda) com teto
    pelo tamanho descompactado declarado no zip, sem descompactar nada."""
    with zipfile.ZipFile(io.BytesIO(dados)) as z:
        for item in z.infolist():
            if item.filename.startswith("xl/worksheets/") or not item.filename.endswith((".xml", ".rels")):
                continue
            if item.file_size > MAX_BYTES_XML:
                raise PlanilhaNaoReconhecida("a planilha é grande demais para o modelo")


def _linhas_xml(dados: bytes, caminho: str):
    """(numero, elemento <row>) da aba, em fluxo: a linha ja' entregue sai da memoria (a aba nunca fica inteira)."""
    with zipfile.ZipFile(io.BytesIO(dados)) as z, z.open(caminho.lstrip("/")) as xml:
        dados_aba, numero = None, 0
        for evento, el in ElementTree.iterparse(xml, events=("start", "end")):
            if evento == "start":
                if el.tag == f"{_NS}sheetData":
                    dados_aba = el
                continue
            if el.tag != f"{_NS}row":
                continue
            numero = int(el.get("r") or numero + 1)
            yield numero, el
            if dados_aba is not None:
                dados_aba.clear()


def _tem_dado(c, textos) -> bool:
    """Celula com valor nao vazio (o que a varredura do openpyxl contaria como dado; " " nao conta)."""
    tipo, v, inline = c.get("t"), c.find(f"{_NS}v"), c.find(f"{_NS}is")
    if tipo == "inlineStr":
        texto = "".join(inline.itertext()) if inline is not None else ""
    elif v is None:
        return False
    elif tipo == "s":
        try:
            texto = str(textos[int(v.text)])
        except (TypeError, ValueError, IndexError):
            return True
    else:
        texto = v.text or ""
    return bool(texto.strip())


def _dado_depois(dados: bytes, caminho: str, ultima: int, textos) -> bool:
    """Ha' valor em alguma linha depois de `ultima`? Formatacao e formula vazia nao contam."""
    for n, linha in _linhas_xml(dados, caminho):
        if n > ultima and any(_tem_dado(c, textos) for c in linha.iter(f"{_NS}c")):
            return True
    return False


def _formulas_sem_valor(dados: bytes, caminho: str, linhas: list[tuple[int, str]],
                        colunas: tuple[int, ...]) -> list[str]:
    """Das linhas lidas sem valor (data_only devolve None), as que tem FORMULA sem nenhum valor guardado: o arquivo
    foi salvo sem calcular (gerado por programa), e a leitura nao pode tratar como vazio. Pelo XML da aba, porque o
    openpyxl devolve None nos dois casos: formula que da' texto vazio (`=SE(...;"";...)`) o Excel grava com
    `<v></v>` e e' vazia de verdade; so' a celula com `<f>` e sem `<v>` ficou sem o calculo."""
    nomes = dict(linhas)
    alvo = {f"{get_column_letter(c + 1)}{n}": n for n in nomes for c in colunas}
    ultima = max(nomes)
    com_formula: list[str] = []
    # uma varredura so', em fluxo, que para na ultima linha procurada (o resto da aba nem e' lido)
    for numero, linha in _linhas_xml(dados, caminho):
        if numero > ultima:
            break
        for el in linha.iter(f"{_NS}c"):
            n = alvo.get(el.get("r", ""))
            if n is not None and el.find(f"{_NS}f") is not None and nomes[n] not in com_formula:
                v = el.find(f"{_NS}v")
                # texto vazio calculado: t="str" com <v></v> (Excel); sem <v>, ou <v> vazio sem tipo (o que o
                # openpyxl e outros geradores gravam), a formula nunca foi calculada
                if v is None or (not (v.text or "") and el.get("t") != "str"):
                    com_formula.append(nomes[n])
    return com_formula


def _recusar_formula_sem_valor(dados: bytes, caminho: str, sem_valor: list[tuple[int, str]], *colunas: int) -> None:
    if not sem_valor:
        return
    nomes = _formulas_sem_valor(dados, caminho, sem_valor, colunas)
    if nomes:
        raise PlanilhaNaoReconhecida(f"{nomes[0]} tem fórmula sem o valor calculado; abra a planilha no Excel, "
                                     "salve e anexe de novo")


def _ler_resumo(wb, dados: bytes) -> list[dict]:
    if ABA not in wb.sheetnames:
        raise PlanilhaNaoReconhecida("falta a aba Resumo")
    ws = wb[ABA]
    caminho = ws._worksheet_path      # o XML da aba, para distinguir formula sem calculo de formula vazia
    with zipfile.ZipFile(io.BytesIO(dados)) as z:
        if z.getinfo(caminho.lstrip("/")).file_size > MAX_BYTES_XML:
            raise PlanilhaNaoReconhecida("a aba Resumo é grande demais para o modelo")
    inicio, col_med, col_per, col_reaj = _colunas(ws)
    out: list[dict] = []
    vistas: dict[str, str] = {}
    sem_valor: list[tuple[int, str]] = []
    # a dimensao declarada nao conta (formatacao ate' a linha 8000 e' comum); o custo fica limitado pelo teto de
    # bytes da aba, pelas colunas lidas e pelo teto de linhas, onde a varredura para. Depois do teto, so' linha com
    # dado recusa (leitura parcial nao serve), conferido direto no XML, sem montar linha nenhuma
    ultima = inicio + MAX_VARRIDAS - 1
    lidas = 0
    for n, linha in enumerate(ws.iter_rows(min_row=inicio, max_row=ultima, max_col=MAX_COLUNAS, values_only=True)):
        lidas = n + 1
        celula = lambda c: linha[c] if c < len(linha) else None   # noqa: E731
        medicao, periodo, valor = celula(col_med), celula(col_per), celula(col_reaj)
        rotulo = _texto(medicao)
        if rotulo.startswith("TOTA"):        # "TOTAIS": acabou a tabela
            break
        if not rotulo and periodo is None:   # linha em branco no meio ou no fim (ou formula sem valor: abaixo)
            if medicao is None:
                sem_valor.append((inicio + n, f"a linha {inicio + n}"))
            continue
        texto_med = str(medicao).strip() if medicao is not None else ""
        nome = texto_med or f"a linha {inicio + n}"     # sem rotulo, a linha da planilha identifica (nunca "None")
        if valor is None or periodo is None:  # medicao sem reajuste previsto (vazia) ou formula sem valor (abaixo)
            sem_valor.append((inicio + n, nome))
            if valor is None:
                continue
        if not isinstance(valor, int | float) or isinstance(valor, bool):
            raise PlanilhaNaoReconhecida(f"o reajuste de {nome} não é um número")
        comp = _competencia(periodo)
        if comp is None:
            _recusar_formula_sem_valor(dados, caminho, sem_valor, col_med, col_per, col_reaj)
            raise PlanilhaNaoReconhecida(f"o período de {nome} não tem data legível")
        if comp in vistas:
            raise PlanilhaNaoReconhecida(f"a competência {comp[5:]}/{comp[:4]} aparece em {vistas[comp]} e em {nome}")
        vistas[comp] = nome
        if len(out) >= MAX_COMPETENCIAS:
            raise PlanilhaNaoReconhecida("competências demais na aba Resumo")
        out.append({"competencia": comp, "valor_centavos": _centavos(valor), "medicao": texto_med,
                    "periodo": periodo.strip() if isinstance(periodo, str) else comp})
    else:
        if lidas >= MAX_VARRIDAS and _dado_depois(dados, caminho, ultima, ws._shared_strings):
            raise PlanilhaNaoReconhecida("a aba Resumo é grande demais para o modelo")
    _recusar_formula_sem_valor(dados, caminho, sem_valor, col_med, col_per, col_reaj)
    if not out:
        raise PlanilhaNaoReconhecida("nenhuma linha de medição com reajuste na aba Resumo")
    return sorted(out, key=lambda x: x["competencia"])
