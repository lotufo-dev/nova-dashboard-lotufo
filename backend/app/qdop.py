"""QDOP - Quadro Demonstrativo de Obras e Projetos (GL.SGI.FRP.DE.250), bloco 8.

Escrito sobre um MODELO recortado do arquivo oficial (cabecalho 1-13, linha 14 como estilo, rodape 16-17),
com VALORES e nao formulas: abre certo em qualquer visualizador. Tudo na cota da Lotufo.

"Faturados" e o reajuste somam as medicoes FATURADAS (nota emitida, protocolada ou recebida), nao so' as
recebidas: conferido contra o QDOP de 22/09/2026, e' o que bate com o oficial (MT-401 24,29 x 24,28 mi; Parque
Atalaia 28,70 x 28,86; 5 Bairros 30,40 x 30,12). So' recebido ficaria muito abaixo, porque o historico carregado
nao marca o recebimento (docs/backlog.md, bloco 8).
"""
import copy
import datetime
import io
import json

import openpyxl

from app import arquivos, contratos, prazos, tempo

COLUNAS = ("seq", "contratante", "tipo", "inicio", "termino", "vigencia", "empreita", "direito", "pi", "va",
           "pi_va", "reajuste_recebido", "faturados", "saldo", "pct_recebido", "pct_a_receber", "status",
           "localidade", "descricao")
EM_REAIS = {"pi", "va", "pi_va", "reajuste_recebido", "faturados", "saldo"}
DATAS = {"inicio", "termino", "vigencia"}
LINHA_DADOS = 14


def _faturado(conn, contrato_id: int, pct, reajuste: bool, mes: str | None = None) -> int:
    """Soma das medicoes faturadas (nota emitida, protocolada, NF emitida do legado, recebida), no atestado
    quando houver (gate Sol do bloco 8: `nota_emitida` e' faturado no vocabulario da diretoria). Com `mes`
    ("AAAA-MM"), so' as que receberam a nota naquele mes (`nota_emitida_em`, data da nota: a mesma regra, fatiada;
    o "Faturado no mes" da Carteira, decisao do Leo de 16/10/2026)."""
    cond = "tipo='reajuste'" if reajuste else "tipo IN ('execucao','projeto')"
    args = (contrato_id,)
    if mes:
        cond += " AND substr(nota_emitida_em,1,7)=?"
        args += (mes,)
    total = conn.execute(f"SELECT COALESCE(SUM(COALESCE(valor_atestado_centavos, valor_protocolado_centavos)), 0) "
                         f"FROM medicoes WHERE contrato_id=? AND estado IN ('nota_emitida','protocolada','nf_emitida','recebida') "
                         f"AND {cond}", args).fetchone()[0]
    return contratos.cota(total, pct) or 0


def _data(iso: str | None) -> datetime.date | None:
    try:
        return datetime.date.fromisoformat(iso[:10]) if iso else None
    except ValueError:
        return None


def linhas(conn, hoje: datetime.date | None = None) -> list[dict]:
    hoje = hoje or tempo.hoje_local()
    out = []
    for c in conn.execute("SELECT c.id, o.regras, o.nome AS orgao_nome, e.papel FROM contratos c "
                          "JOIN orgaos o ON o.id=c.orgao_id LEFT JOIN entidades e ON e.id=c.entidade_id "
                          "WHERE c.estado<>'rescindido' ORDER BY c.execucao_inicio, c.id").fetchall():
        v = contratos.vigente(conn, c["id"])
        regras = json.loads(c["regras"] or "{}")
        pct = v["percentual_direito_bp"]
        pi = contratos.cota(v["valor_original_centavos"], pct) or 0
        pi_va = contratos.cota(v["valor_vigente_centavos"], pct) or 0
        faturados = _faturado(conn, c["id"], pct, False)
        s = prazos.situacao(conn, c["id"], hoje)
        status = ("CONCLUÍDA" if v.get("encerramento_iniciado_em") or v["estado"] == "encerrado"
                  else "PARALISADA" if s["paralisacao"] or v["estado"] == "suspenso" else "EM ANDAMENTO")
        out.append({"contrato_id": c["id"], "obra": v.get("nome_curto"),
                    "contratante": regras.get("nome_qdop") or c["orgao_nome"],
                    "tipo": regras.get("tipo_contratante") or "PÚBLICO",
                    "inicio": v.get("execucao_inicio"), "termino": v.get("execucao_fim_vigente"),
                    "vigencia": v.get("vigencia_fim_vigente"),
                    "empreita": "PRÓPRIA" if (c["papel"] or "lotufo") == "lotufo" else "CONSÓRCIO",
                    "direito": (pct or contratos.DIREITO_INTEGRAL_BP) / contratos.DIREITO_INTEGRAL_BP,
                    "pi": pi, "va": pi_va - pi, "pi_va": pi_va,
                    "reajuste_recebido": _faturado(conn, c["id"], pct, True), "faturados": faturados,
                    "saldo": pi_va - faturados, "pct_recebido": faturados / pi_va if pi_va else 0,
                    "pct_a_receber": (pi_va - faturados) / pi_va if pi_va else 0, "status": status,
                    "localidade": f"{(v.get('municipio') or '').upper()} - MT" if v.get("municipio") else "",
                    "descricao": v.get("objeto_completo") or v.get("objeto") or ""})
    for i, linha in enumerate(out, 1):
        linha["seq"] = i
    return out


def _modelo(conn):
    r = conn.execute("SELECT valor FROM cadastro_geral WHERE chave='qdop'").fetchone()
    if r is None:
        raise ValueError("Modelo do QDOP não carregado: rode `cadastro aplicar` com geral.yaml (qdop.modelo)")
    return arquivos.caminho_seguro(json.loads(r["valor"])["modelo"])


def gerar_xlsx(conn, hoje: datetime.date | None = None) -> bytes:
    hoje = hoje or tempo.hoje_local()
    ls = linhas(conn, hoje)
    wb = openpyxl.load_workbook(_modelo(conn))
    ws = wb.active
    estilo = {col: copy.copy(ws.cell(LINHA_DADOS, col)._style) for col in range(1, len(COLUNAS) + 1)}
    altura = ws.row_dimensions[LINHA_DADOS].height
    rodape = [([(cel.column, cel.value, copy.copy(cel._style)) for cel in ws[r] if cel.value], ws.row_dimensions[r].height)
              for r in (16, 17)]
    ws.delete_rows(LINHA_DADOS, ws.max_row - LINHA_DADOS + 1)
    for i, linha in enumerate(ls):
        r = LINHA_DADOS + i
        for col, chave in enumerate(COLUNAS, 1):
            v = linha[chave]
            if chave in EM_REAIS:
                v = v / 100
            elif chave in DATAS:
                v = _data(v)
            ws.cell(r, col, v)._style = copy.copy(estilo[col])
        ws.row_dimensions[r].height = altura
    ultima = LINHA_DADOS - 1 + len(ls)
    for linha_tot, grupo in ((9, "PRÓPRIA"), (10, "CONSÓRCIO"), (11, None)):
        sel = [x for x in ls if grupo is None or x["empreita"] == grupo]
        for col, chave in zip("IJKLMN", ("pi", "va", "pi_va", "reajuste_recebido", "faturados", "saldo")):
            ws[f"{col}{linha_tot}"] = sum(x[chave] for x in sel) / 100
    ws["A9"] = len(ls)
    propria = [x for x in ls if x["empreita"] == "PRÓPRIA"]      # O9/P9 do oficial: =M9/(I9+J9), so' a PRÓPRIA
    base = sum(x["pi_va"] for x in propria)
    ws["O9"] = sum(x["faturados"] for x in propria) / base if base else 0
    ws["P9"] = sum(x["saldo"] for x in propria) / base if base else 0
    ws["S7"] = hoje
    for k, (linha, alt) in enumerate(rodape):
        for col, valor, est in linha:
            ws.cell(ultima + 2 + k, col, valor)._style = est
        ws.row_dimensions[ultima + 2 + k].height = alt
    ws.print_area = f"A1:S{ultima + 3}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
