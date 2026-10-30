"""FRC gerada pela dashboard (2026-10-09): a FRC do consorcio do cadastro de exemplo (Obra Consorcio, 14a medicao
de execucao) reproduzida a partir das quatro notas, checklist que bloqueia, numeracao e a FRC assinada no envio.

As notas sao textos de fixture (uma da empresa propria pelo portal de Cuiaba, outra da consorciada Beta pelo
Centi) com numero, municipio, CNPJ e valores trocados pelos do espelho de exemplo (R$ 10.000.000,07 cheio:
Cuiaba 60% e Varzea Grande 40%, metade para cada empresa). INSS/ISS/IRRF seguem a regra do espelho (11% e 5%
sobre 40%, 1,2% sobre o total). O cadastro da FRC do consorcio de exemplo e' semeado pelo proprio teste: em
producao ele mora na migracao 032."""
import io
import json

import openpyxl
import pypdf
import pytest

from app import cadastro_aplicar, cadastro_arquivos, espelho, espelhos, frc, leitura_pdf, medicoes, notas, notas_envio, protocolo
from tests.conftest import CADASTRO_EXEMPLO, _pdf_com_texto
from tests.test_notas import CNPJ_BETA, CNPJ_PROPRIA, _fixture, _texto, conn, usuario  # noqa: F401  (fixtures)

LOTUFO, BETA = CNPJ_PROPRIA, CNPJ_BETA
CONSORCIO = "consorcio-exemplo"
ULTIMA = 40                      # a ultima FRC emitida antes do sistema: a primeira gerada sai 41
# (empresa, numero, municipio, bruto, inss, iss, irrf), em centavos. A ultima e' do leiaute de Cuiaba. O cheio
# (R$ 10.000.000,07) tem centavos: a metade de Varzea Grande cai em .5 e o ISS/INSS/IRRF truncam.
NOTAS_FRC = [
    ("beta", 20260000204, "Várzea Grande", 200000001, 8800000, 4000000, 2400000),
    ("beta", 20260000205, "Cuiabá", 300000002, 13200000, 6000000, 3600000),
    ("lotufo", 2284, "Cuiabá", 300000002, 13200000, 6000000, 3600000),
    ("lotufo", 2286, "Várzea Grande", 200000002, 8800000, 4000000, 2400000),
]
DADOS_FRC = {
    "logo": None,
    "prestador": {"razao_social": "CONSÓRCIO EXEMPLO", "endereco": "RUA DE EXEMPLO, 100, SALA 1, BAIRRO EXEMPLO,",
                  "municipio": "CUIABÁ - MT / BRASIL", "cep": "00000-000", "cnpj": "31.200.303/0001-81",
                  "inscricao_municipal": "123456", "telefone": "(00) 0000-0000", "email": "contato@exemplo.com.br"},
    "tomador": {"razao_social": "SECRETARIA DE EXEMPLO", "endereco": "AV. DE EXEMPLO", "municipio": "CUIABÁ - MT / BRASIL",
                "cep": "00000-001", "cnpj": "31.300.101/0001-01", "inscricao_municipal": "654321",
                "telefone": "(00) 0000-0001", "email": "contato@exemplo.gov.br"},
    "descricao": {
        "execucao": "{n}ª MEDIÇÃO DE CONTRATO - PERÍODO DE {inicio} A {fim} - OBRAS DE EXEMPLO, CONFORME IC 002/2025.",
        "reajuste": "REAJUSTE DA {n}ª MEDIÇÃO DE CONTRATO - PERÍODO DE {inicio} A {fim} - OBRAS DE EXEMPLO."},
    "rotulo_empresa": "RAZÃO SOCIAL",
    "consorciadas": [{"chave": "lotufo", "cnpj": "31.200.101/0001-30", "razao_social": "CONSTRUTORA EXEMPLO LTDA"},
                     {"chave": "beta", "cnpj": "31.200.202/0001-00", "razao_social": "ENGENHARIA BETA S.A."}],
}


def _semear_frc(conn, ultimo=ULTIMA, **dados):
    """O cadastro da FRC do Consorcio Exemplo (formulario A) gravado direto, sem logo, por cima do frc.yaml."""
    conn.execute("INSERT OR REPLACE INTO frc_consorcios (consorcio, formulario, ultimo_anterior, dados) "
                 "VALUES (?, 'A', ?, ?)", (CONSORCIO, ultimo, json.dumps({**DADOS_FRC, **dados}, ensure_ascii=False)))
    conn.commit()


def _cuiaba(cnpj, numero, municipio, bruto, inss, iss, irrf) -> str:
    f = espelho._fmt
    ret = inss + iss + irrf
    t = _fixture("3102-integra-vg.txt").replace(LOTUFO, cnpj)
    t = t.replace("Número da Nota Fiscal\n3102", f"Número da Nota Fiscal\n{numero}").replace("Várzea Grande", municipio)
    t = t.replace("12ª MEDIÇÃO", "14ª MEDIÇÃO").replace("PERÍODO DE 01/06/2026 A 30/06/2026", "PERÍODO DE 01/08/2026 A 31/08/2026")
    t = t.replace("R$846.788,31", "R$" + f(bruto)).replace("Vl. ISSQN: R$16.935,77", "Vl. ISSQN: R$" + f(iss))
    t = t.replace("Vl. IRRF: R$10.161,45", "Vl. IRRF: R$" + f(irrf)).replace("Vl. CP Retido: R$37.258,68", "Vl. CP Retido: R$" + f(inss))
    return t.replace("R$64.355,90 R$2.876,20 R$319,58 R$782.432,41 R$782.432,41",
                     f"R${f(ret)} R$2.876,20 R$319,58 R${f(bruto - ret)} R${f(bruto - ret)}")


def _centi(numero, municipio, bruto, inss, iss, irrf) -> str:
    f = espelho._fmt
    t = _fixture("2026812-integra-beta-vg.txt").replace("20260000812", str(numero))   # tambem na chave de acesso
    t = t.replace("VÁRZEA GRANDE - MT", municipio.upper() + " - MT").replace("11ª MEDIÇÃO", "14ª MEDIÇÃO")
    t = t.replace("PERÍODO DE 01/05/2026 A 31/05/2026", "PERÍODO DE 01/08/2026 A 31/08/2026")
    for de, para in (("914.477,97", bruto), ("40.237,02", inss), ("10.973,73", irrf), ("51.210,75", inss + irrf),
                     ("18.289,56", iss), ("844.977,66", bruto - inss - irrf - iss)):
        t = t.replace(de, f(para))
    return t


def _texto_da(nota) -> str:
    empresa, numero, municipio, *v = nota
    if empresa == "beta":
        return _centi(numero, municipio, *v)
    return _cuiaba(LOTUFO, numero, municipio, *v)


def _soma(i: int) -> int:
    return sum(n[i] for n in NOTAS_FRC)


@pytest.fixture
def consorcio14(conn, usuario):  # noqa: F811
    """Cadastro de exemplo, FRC do consorcio semeada, espelho da 14a com o cheio de R$ 10.000.000,07."""
    cadastro_aplicar.aplicar(conn, 1, cadastro_arquivos.preparar(CADASTRO_EXEMPLO), ensaio=False)
    _semear_frc(conn)
    return _medicao14(conn, usuario)


def _medicao14(conn, usuario):  # noqa: F811
    cid = conn.execute("SELECT id FROM contratos WHERE nome_curto='Obra Consórcio'").fetchone()[0]
    eid = espelhos.gerar(conn, usuario, cid, 14, "2026-08", {"medicao": _soma(3)})
    mid = conn.execute("SELECT m.id FROM espelho_medicoes em JOIN medicoes m ON m.id=em.medicao_id "
                       "WHERE em.espelho_id=? AND m.tipo='execucao'", (eid,)).fetchone()[0]
    medicoes.marcar_espelho_enviado(conn, 1, mid)
    return cid, mid


def _anexar(conn, usuario, mid, _texto, lista):  # noqa: F811
    out = []
    for i, nota in enumerate(lista):
        _texto(_texto_da(nota))
        out.append(notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto(f"nota {nota[1]}"), f"NFSE {i}.pdf"))
    return out


def test_frc_reproduzida_das_quatro_notas(conn, usuario, consorcio14, _texto):  # noqa: F811
    _, mid = consorcio14
    s = frc.situacao(conn, mid)
    assert s["disponivel"] and (s["ultima"], s["sugerido"]) == (ULTIMA, ULTIMA + 1) and s["frc"] is None
    assert [(i["empresa"], i["municipio"]) for i in s["itens"]] == [
        ("Exemplo", "Cuiabá"), ("Beta", "Cuiabá"), ("Exemplo", "Várzea Grande"), ("Beta", "Várzea Grande")]
    r = _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    assert [x["frc_gerada"] for x in r] == [None] * 3 + [ULTIMA + 1]      # a ultima nota gera sozinha, com ultima + 1
    s = frc.situacao(conn, mid)
    assert s["completa"] and s["bloqueio"] is None and s["frc"]["numero"] == ULTIMA + 1 and s["ultima"] == ULTIMA + 1
    assert all(i["situacao"] == "anexada" for i in s["itens"])
    assert {i["numero"] for i in s["itens"] if i["empresa"] == "Beta"} == {"2026204", "2026205"}
    f = conn.execute("SELECT * FROM frcs WHERE medicao_id=?", (mid,)).fetchone()
    bruto, inss, iss, irrf = _soma(3), _soma(4), _soma(5), _soma(6)
    liquido = bruto - inss - iss - irrf
    assert (f["bruto_centavos"], f["inss_centavos"], f["irrf_centavos"], f["iss_centavos"], f["liquido_centavos"]) == (
        1000000007, 44000000, 12000000, 20000000, 924000007) == (bruto, inss, irrf, iss, liquido)
    linhas = json.loads(f["notas"])
    assert [x["numero"] for x in linhas] == ["2284", "2286", "2026204", "2026205"]      # por consorciada, depois pelo numero
    assert [(x["bruto"], x["retencoes"], x["liquido"]) for x in linhas][-2:] == [(200000001, 15200000, 184800001),
                                                                                (300000002, 22800000, 277200002)]
    texto = frc.arquivo(conn, mid, "pdf").read_bytes()
    t = pypdf.PdfReader(io.BytesIO(texto)).pages[0].extract_text()
    for trecho in ("FATURA DE RECEBIMENTO DE CONSÓRCIO", f"{ULTIMA + 1:05d}", "GL.SGI.FRP.300", "CONSÓRCIO EXEMPLO",
                   "31.200.303/0001-81", "14ª MEDIÇÃO DE CONTRATO - PERÍODO DE 01/08/2026 A 31/08/2026",
                   "10.000.000,07", "440.000,00", "120.000,00", "200.000,00", "9.240.000,07", "2026204",
                   "ENGENHARIA BETA S.A.", "contato@exemplo.gov.br"):
        assert trecho in t, trecho
    wb = openpyxl.load_workbook(frc.arquivo(conn, mid, "xlsx"))
    valores = [c.value for linha in wb.active.iter_rows() for c in linha if c.value is not None]
    assert 9240000.07 in valores and f"{ULTIMA + 1:05d}" in valores and "2026204" in valores
    assert frc.nome_arquivo(conn, mid, "pdf") == f"FRC {ULTIMA + 1} - CONSÓRCIO EXEMPLO (14ª MEDIÇÃO).pdf"


def test_frc_do_cadastro_sai_com_o_logo_e_o_tomador_do_cadastro(conn, usuario, _texto):  # noqa: F811
    """Issue #110: sem semente da migracao, a FRC vem do frc.yaml (cadastro-exemplo). O logo, copiado para a pasta
    de dados pelo `cadastro aplicar`, vai no PDF e na planilha; o e-mail do tomador e' o do cadastro."""
    cadastro_aplicar.aplicar(conn, 1, cadastro_arquivos.preparar(CADASTRO_EXEMPLO), ensaio=False)
    _, mid = _medicao14(conn, usuario)
    assert _anexar(conn, usuario, mid, _texto, NOTAS_FRC)[-1]["frc_gerada"] == ULTIMA + 1
    pagina = pypdf.PdfReader(io.BytesIO(frc.arquivo(conn, mid, "pdf").read_bytes())).pages[0]
    assert len(pagina.images) == 1 and "fiscal@exemplo.gov.br" in pagina.extract_text()
    assert len(openpyxl.load_workbook(frc.arquivo(conn, mid, "xlsx")).active._images) == 1


def test_frc_com_logo_que_nao_esta_na_pasta_de_dados_e_recusada(conn, usuario, _texto):  # noqa: F811
    """Deploy sem `cadastro aplicar` (issue #110): a linha antiga aponta para o logo que saiu do app. A FRC nao sai
    sem o papel timbrado: a geracao automatica nao gera e a manual explica o que falta."""
    cadastro_aplicar.aplicar(conn, 1, cadastro_arquivos.preparar(CADASTRO_EXEMPLO), ensaio=False)
    dados = json.loads(conn.execute("SELECT dados FROM frc_consorcios").fetchone()[0])
    conn.execute("UPDATE frc_consorcios SET dados=?", (json.dumps({**dados, "logo": "frc-integra-brt.jpg"}),))
    conn.commit()
    _, mid = _medicao14(conn, usuario)
    assert _anexar(conn, usuario, mid, _texto, NOTAS_FRC)[-1]["frc_gerada"] is None
    with pytest.raises(ValueError, match="Logo da FRC não encontrado na pasta de dados: aplique o cadastro"):
        frc.gerar(conn, usuario, mid, ULTIMA + 1)
    assert conn.execute("SELECT COUNT(*) FROM frcs").fetchone()[0] == 0


def test_natureza_sem_descricao_no_cadastro_recusa_com_mensagem(conn, usuario, _texto):  # noqa: F811
    """O frc.yaml pode trazer so' as naturezas que o consorcio usa (Complexo Leblon nao tem projeto): medicao de
    natureza sem descricao nao gera, com a mensagem, sem estourar."""
    cadastro_aplicar.aplicar(conn, 1, cadastro_arquivos.preparar(CADASTRO_EXEMPLO), ensaio=False)
    _semear_frc(conn, descricao={"reajuste": "REAJUSTE DA {n}ª MEDIÇÃO - {inicio} A {fim}"})
    _, mid = _medicao14(conn, usuario)
    assert _anexar(conn, usuario, mid, _texto, NOTAS_FRC)[-1]["frc_gerada"] is None
    with pytest.raises(ValueError, match="Cadastro da FRC sem descrição para 14ª medição"):
        frc.gerar(conn, usuario, mid, ULTIMA + 1)


def test_checklist_com_nota_faltando_bloqueia(conn, usuario, consorcio14, _texto):  # noqa: F811
    _, mid = consorcio14
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC[:3])
    s = frc.situacao(conn, mid)
    assert not s["completa"] and s["frc"] is None
    assert s["bloqueio"] == "Falta 1 nota: Exemplo · Várzea Grande"
    with pytest.raises(ValueError, match="Exemplo"):
        frc.gerar(conn, usuario, mid, ULTIMA + 1)


def test_nota_sem_leitura_confiavel_bloqueia_e_rascunho_nao_conta(conn, usuario, consorcio14, _texto):  # noqa: F811
    _, mid = consorcio14
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC[:3])
    # rascunho da ultima nota: mesmo nome de arquivo, sem numero no conteudo -> nao atende a saida
    _texto(_texto_da(NOTAS_FRC[3]).replace("Número da Nota Fiscal\n2286", "Número da Nota Fiscal\n0"))
    notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto("rascunho"), "NFSE 2286.pdf")
    s = frc.situacao(conn, mid)
    assert [x["motivo"] for x in s["nao_lidas"]] == ["sem-numero"] and not s["completa"]
    # a nota de verdade, com o liquido impresso que nao fecha com bruto - retencoes
    _texto(_texto_da(NOTAS_FRC[3]).replace("R$1.848.000,02 R$1.848.000,02", "R$1.848.000,03 R$1.848.000,03"))
    r = notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto("2286 errada"), "NFSE 2286.pdf")
    assert r["frc_gerada"] is None
    s = frc.situacao(conn, mid)
    assert s["completa"] and s["frc"] is None
    assert s["bloqueio"].startswith("Nota sem leitura confiável — NF 2286 (Exemplo)")
    with pytest.raises(ValueError, match="sem leitura confiável"):
        frc.gerar(conn, usuario, mid, ULTIMA + 1)


def test_arrastar_de_novo_rele_a_nota_que_nao_tinha_sido_lida(conn, usuario, consorcio14, _texto):  # noqa: F811
    # 09/10/2026: as notas novas da consorciada chegaram antes do leitor conhecer o modelo; a pessoa arrasta o
    # MESMO PDF de novo depois da correcao e a nota passa a contar (e a FRC sai sozinha)
    _, mid = consorcio14
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC[:3])
    _texto("NOTA FISCAL DE SERVIÇOS ELETRÔNICA\nleiaute que o leitor ainda nao conhece")
    a = notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto("nota 2286"), "NFSE 2286.pdf")
    assert not a["leitura"]["ok"] and a["frc_gerada"] is None and not frc.situacao(conn, mid)["completa"]
    _texto(_texto_da(NOTAS_FRC[3]))
    b = notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto("nota 2286"), "NFSE 2286.pdf")
    assert b["repetido"] and b["documento_id"] == a["documento_id"]
    assert b["frc_gerada"] == ULTIMA + 1 and frc.situacao(conn, mid)["completa"]


def test_numero_sugerido_e_ultima_registrada(conn, usuario, consorcio14, _texto):  # noqa: F811
    cid, mid = consorcio14
    n = ULTIMA + 1
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    assert frc.ultima(conn, CONSORCIO) == n
    frc.gerar(conn, usuario, mid, n + 8)                  # numero digitado: gera de novo e a ultima acompanha
    s = frc.situacao(conn, mid)
    assert (s["frc"]["numero"], s["ultima"], s["sugerido"]) == (n + 8, n + 8, n + 8)
    assert conn.execute("SELECT COUNT(*) FROM frcs WHERE medicao_id=? AND substituida_em IS NOT NULL",
                        (mid,)).fetchone()[0] == 1
    frc.gerar(conn, usuario, mid, n)                      # volta ao n: a n + 8 foi substituida
    assert frc.ultima(conn, CONSORCIO) == n
    outra = medicoes.criar(conn, 1, {"contrato_id": cid, "numero": 15, "tipo": "execucao", "competencia": "2026-09"})
    assert frc.situacao(conn, outra)["sugerido"] == n + 1
    # o mesmo numero em outra medicao do consorcio e' recusado (aqui a outra nem tem as notas: o numero vem antes)
    conn.execute("INSERT INTO frcs (medicao_id, consorcio, numero, data_geracao, bruto_centavos, inss_centavos, "
                 "irrf_centavos, iss_centavos, liquido_centavos, notas, pdf, pdf_sha256, xlsx, xlsx_sha256) "
                 "VALUES (?, ?, ?, '2026-10-09', 0, 0, 0, 0, 0, '[]', 'x', 'x', 'x', 'x')", (outra, CONSORCIO, n + 1))
    conn.commit()
    with pytest.raises(ValueError, match=f"FRC nº {n + 1} já é a da 15ª Medição"):
        frc.gerar(conn, usuario, mid, n + 1)


def test_frc_gerada_nao_satisfaz_o_envio_e_a_assinada_sim(conn, usuario, consorcio14, _texto):  # noqa: F811
    cid, mid = consorcio14
    n = ULTIMA + 1
    c = protocolo._contrato(conn, cid)
    m = dict(conn.execute("SELECT * FROM medicoes WHERE id=?", (mid,)).fetchone())
    # issue #148: sem FRC gerada, a etapa 4 manda gerar na etapa 3 e nao oferece baixar nada
    itens, _, _ = notas_envio._itens(conn, c, "sinfra", [m], __import__("datetime").date(2026, 10, 9))
    item = next(i for i in itens if i["chave"] == f"frc:{mid}")
    assert item["frc_gerada"] is None and item["detalhe"].startswith("Falta a FRC assinada: gere na etapa 3")
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    assert frc.situacao(conn, mid)["frc"] is not None
    assert protocolo.faltas_consorcio(conn, c, [m]) == ["14ª Medição: FRC assinada não anexada"]
    itens, _, _ = notas_envio._itens(conn, c, "sinfra", [m], __import__("datetime").date(2026, 10, 9))
    item = next(i for i in itens if i["chave"] == f"frc:{mid}")
    assert (item["estado"], item["acao"]) == ("falta", "anexar_frc")
    # issue #148: a etapa 4 diz que falta a assinada e oferece baixar a gerada
    assert item["frc_gerada"] == n
    assert item["detalhe"] == "Falta a assinada: baixe a FRC gerada, assine no DocuSign e anexe aqui"
    assert next(i for i in itens if i["chave"] == "consorcio")["estado"] == "ok"
    # FRC desatualizada (as notas mudaram depois dela): nao se oferece para assinar, manda gerar de novo
    linhas = conn.execute("SELECT notas FROM frcs WHERE medicao_id=? AND substituida_em IS NULL", (mid,)).fetchone()[0]
    conn.execute("UPDATE frcs SET notas='[]' WHERE medicao_id=? AND substituida_em IS NULL", (mid,))
    conn.commit()
    itens, _, _ = notas_envio._itens(conn, c, "sinfra", [m], __import__("datetime").date(2026, 10, 9))
    item = next(i for i in itens if i["chave"] == f"frc:{mid}")
    assert item["frc_gerada"] is None and "gere de novo na etapa 3" in item["detalhe"]
    conn.execute("UPDATE frcs SET notas=? WHERE medicao_id=? AND substituida_em IS NULL", (linhas, mid))
    conn.commit()
    _texto(_texto_da_frc(conn, mid))
    notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto("assinada"), f"FRC {n} assinada.pdf", tipo_doc="fatura")
    assert protocolo.faltas_consorcio(conn, c, [m]) == []
    itens, _, _ = notas_envio._itens(conn, c, "sinfra", [m], __import__("datetime").date(2026, 10, 9))
    item = next(i for i in itens if i["chave"] == f"frc:{mid}")
    assert (item["estado"], item["detalhe"]) == ("ok", f"Assinada anexada · FRC {n} assinada.pdf")
    assert frc.situacao(conn, mid)["assinada"] is True


def _item_frc(conn, cid, mid):
    c = protocolo._contrato(conn, cid)
    m = dict(conn.execute("SELECT * FROM medicoes WHERE id=?", (mid,)).fetchone())
    itens, _, _ = notas_envio._itens(conn, c, "sinfra", [m], __import__("datetime").date(2026, 10, 9))
    return next(i for i in itens if i["chave"] == f"frc:{mid}"), protocolo.faltas_consorcio(conn, c, [m])


def _texto_da_frc(conn, mid) -> str:
    """O texto do PDF da FRC vigente, como o DocuSign o devolve assinado (o texto e' o mesmo)."""
    p = frc.arquivo(conn, mid, "pdf")
    return pypdf.PdfReader(io.BytesIO(p.read_bytes())).pages[0].extract_text() if p else "qualquer coisa"


def _assinar(conn, usuario, mid, _texto, nome, texto=None):  # noqa: F811
    _texto(texto if texto is not None else _texto_da_frc(conn, mid))
    return notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto(nome), nome, tipo_doc="fatura")


def test_assinada_de_frc_substituida_nao_deixa_o_item_pronto(conn, usuario, consorcio14, _texto):  # noqa: F811
    """Issue #169: a assinada fica ligada a' FRC que ela assina; gerada de novo (mesmo numero), a anexada nao vale."""
    cid, mid = consorcio14
    n = ULTIMA + 1
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    primeira = conn.execute("SELECT id FROM frcs WHERE medicao_id=? AND substituida_em IS NULL", (mid,)).fetchone()[0]
    velha = _assinar(conn, usuario, mid, _texto, f"FRC {n} assinada.pdf")["documento_id"]
    assert conn.execute("SELECT frc_id FROM medicao_documentos WHERE id=?", (velha,)).fetchone()[0] == primeira
    assert _item_frc(conn, cid, mid)[0]["estado"] == "ok"
    frc.gerar(conn, usuario, mid, n)                       # notas trocadas, FRC gerada de novo: a assinada e' da velha
    item, faltas = _item_frc(conn, cid, mid)
    recusa = f"A FRC assinada anexada não é a da FRC nº {n} vigente: baixe a FRC gerada, assine no DocuSign e anexe aqui"
    assert (item["estado"], item["acao"], item["detalhe"], item["frc_gerada"]) == ("falta", "anexar_frc", recusa, n)
    assert faltas == []                                    # nao e' "nao anexada": e' bloqueio (protocolo.frc_recusadas)
    s = frc.situacao(conn, mid)
    assert (s["assinada"], s["assinada_recusa"]) == (False, recusa)
    a = frc.assinada(conn, mid)
    assert (a["controla"], a["documento"], a["recusa"], a["frc"]["numero"]) == (True, None, recusa, n)
    assert item["bloqueia"] is True
    assert protocolo._docs_nota(conn, [mid])[-1]["tipo_doc"] == "nfs"      # a velha nao vai no envio
    nova = _assinar(conn, usuario, mid, _texto, f"FRC {n} assinada de novo.pdf")["documento_id"]
    item, faltas = _item_frc(conn, cid, mid)
    assert (item["estado"], item["detalhe"], faltas) == ("ok", f"Assinada anexada · FRC {n} assinada de novo.pdf", [])
    assert [d["id"] for d in protocolo._docs_nota(conn, [mid]) if d["tipo_doc"] == "fatura"] == [nova]


def test_assinada_so_se_anexa_sobre_a_frc_vigente_atualizada(conn, usuario, consorcio14, _texto):  # noqa: F811
    cid, mid = consorcio14
    n = ULTIMA + 1
    with pytest.raises(ValueError, match="Gere a FRC na etapa 3"):
        _assinar(conn, usuario, mid, _texto, "FRC sem gerar.pdf")
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    _assinar(conn, usuario, mid, _texto, f"FRC {n} assinada.pdf")
    # as notas mudaram depois da FRC: a assinada dela deixa de valer e outra nao se anexa ate' gerar de novo
    conn.execute("UPDATE frcs SET notas='[]' WHERE medicao_id=? AND substituida_em IS NULL", (mid,))
    conn.commit()
    item, faltas = _item_frc(conn, cid, mid)
    assert (item["estado"], item["frc_gerada"], faltas) == ("falta", None, [])
    assert item["detalhe"] == (f"As notas mudaram depois da FRC nº {n}: gere de novo na etapa 3 (Nota fiscal), "
                               "assine no DocuSign e anexe aqui")
    with pytest.raises(ValueError, match=f"As notas mudaram depois da FRC nº {n}"):
        _assinar(conn, usuario, mid, _texto, "FRC outra.pdf")
    assert conn.execute("SELECT COUNT(*) FROM medicao_documentos WHERE medicao_id=? AND tipo_doc='fatura'",
                        (mid,)).fetchone()[0] == 1


def test_assinada_sem_vinculo_se_religa_ao_arrastar_de_novo_e_pdf_sem_texto_e_recusado(conn, usuario, consorcio14,
                                                                                         _texto):  # noqa: F811
    """Revisor adversarial na PR #179 (aviso 2): a assinada de antes da migracao 038 (sem vinculo) que assina a FRC
    vigente volta a valer quando o mesmo PDF e' arrastado de novo; PDF sem texto nao se confere e e' recusado."""
    cid, mid = consorcio14
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    doc = _assinar(conn, usuario, mid, _texto, "FRC assinada.pdf")["documento_id"]
    conn.execute("UPDATE medicao_documentos SET frc_id=NULL WHERE id=?", (doc,))       # anexada antes da 038
    conn.commit()
    assert _item_frc(conn, cid, mid)[0]["estado"] == "falta"
    r = _assinar(conn, usuario, mid, _texto, "FRC assinada.pdf")                       # o mesmo arquivo
    assert (r["documento_id"], r["repetido"]) == (doc, True)
    assert _item_frc(conn, cid, mid)[0]["estado"] == "ok"
    with pytest.raises(ValueError, match="não tem texto para conferir"):
        _assinar(conn, usuario, mid, _texto, "FRC digitalizada.pdf", texto="")


def test_conferencia_da_assinada_por_palavra_inteira_e_falha_de_leitura_com_a_frase_dela(conn, usuario, consorcio14,
                                                                                          _texto, monkeypatch):  # noqa: F811
    """Revisor adversarial na PR #179 (rodada 2): o numero da nota dentro de outro numero (2284 em 122845) nao confere;
    leitor ocupado ou lento nao vira "digitalizacao": sobe a frase da leitura, que manda tentar de novo."""
    _, mid = consorcio14
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    certo = _texto_da_frc(conn, mid)
    assert "\n2284 " in certo
    with pytest.raises(ValueError, match="Este PDF não é a FRC nº"):
        _assinar(conn, usuario, mid, _texto, "FRC pedaco.pdf", texto=certo.replace("\n2284 ", "\n122845 "))
    def ocupado(pdf):
        raise leitura_pdf.TextoInadequado(leitura_pdf._LOTADO)
    monkeypatch.setattr(notas.leitura_pdf, "texto_de_pdf_com_limite", ocupado)
    with pytest.raises(leitura_pdf.TextoInadequado) as erro:
        notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto("ocupado"), "FRC.pdf", tipo_doc="fatura")
    assert str(erro.value) == leitura_pdf._LOTADO and "não tem texto" not in str(erro.value)
    # quebras de linha do DocuSign no lugar dos espacos continuam conferindo
    assert _assinar(conn, usuario, mid, _texto, "FRC docusign.pdf", texto=certo.replace(" ", "\n"))["documento_id"]


def test_numero_e_valor_quebrados_pelo_docusign_conferem_sem_aceitar_pedaco(conn, usuario, consorcio14, _texto):  # noqa: F811
    """Rodada 3 da PR #179 (S2): o DocuSign quebra numero e valor na coluna estreita ("10.000.\\n000,07", "22 84");
    a conferencia junta os pedacos, mas segue por palavra inteira: "122 84" nao e' a nota 2284."""
    _, mid = consorcio14
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    certo = _texto_da_frc(conn, mid)
    bruto = espelho._fmt(_soma(3))
    assert bruto == "10.000.000,07" and bruto in certo and "\n2284 " in certo and "00041" in certo
    quebrado = certo.replace(bruto, "10.000.\n000,0 7").replace("\n2284 ", "\n22 84 ").replace("00041", "000\n41")
    with pytest.raises(ValueError, match="Este PDF não é a FRC nº"):
        _assinar(conn, usuario, mid, _texto, "FRC pedaco.pdf", texto=quebrado.replace("\n22 84 ", "\n122 84 "))
    with pytest.raises(ValueError, match="Este PDF não é a FRC nº"):
        _assinar(conn, usuario, mid, _texto, "FRC colada.pdf", texto=quebrado.replace("\n22 84 ", "\n22 84,5 "))
    assert _assinar(conn, usuario, mid, _texto, "FRC docusign quebrada.pdf", texto=quebrado)["documento_id"]


def test_assinada_conferida_dentro_da_transacao_do_anexo(conn, usuario, consorcio14, _texto, monkeypatch):  # noqa: F811
    """Rodada 3 da PR #179 (S3): a FRC fica desatualizada entre a leitura do PDF e a gravacao (outra pessoa troca
    uma nota); a conferencia roda sob o BEGIN IMMEDIATE, depois disso, e recusa: nada fica ligado a' versao velha."""
    _, mid = consorcio14
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    gravar = notas.arquivos.gravar_bytes

    def no_meio(*a, **k):                  # roda entre a leitura do texto (fora) e a transacao do anexo
        conn.execute("UPDATE frcs SET notas='[]' WHERE medicao_id=? AND substituida_em IS NULL", (mid,))
        conn.commit()
        return gravar(*a, **k)
    monkeypatch.setattr(notas.arquivos, "gravar_bytes", no_meio)
    with pytest.raises(ValueError, match="As notas mudaram depois da FRC nº"):
        _assinar(conn, usuario, mid, _texto, "FRC assinada.pdf")
    assert conn.execute("SELECT COUNT(*) FROM medicao_documentos WHERE medicao_id=? AND tipo_doc='fatura'",
                        (mid,)).fetchone()[0] == 0


def test_pdf_corrompido_da_assinada_vira_frase_em_portugues(conn, usuario, consorcio14, _texto, monkeypatch):  # noqa: F811
    """Rodada 3 da PR #179 (aviso): PDF que o leitor nao abre nao mostra "PDF inválido (PdfReadError): ..."; a falha
    de fila (leitor ocupado) segue com a frase dela, que manda tentar de novo."""
    _, mid = consorcio14
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC)

    def corrompido(pdf):                   # o que leitura_pdf.com_limite levanta quando o pypdf nao abre o arquivo
        raise ValueError("PDF inválido (PdfReadError): EOF marker not found")
    monkeypatch.setattr(notas.leitura_pdf, "texto_de_pdf_com_limite", corrompido)
    with pytest.raises(ValueError) as erro:
        notas.anexar_pdf(conn, usuario, mid, b"%PDF-1.4\nquebrado", "FRC quebrada.pdf", tipo_doc="fatura")
    assert str(erro.value) == frc.CORROMPIDO and not isinstance(erro.value, leitura_pdf.TextoInadequado)
    assert conn.execute("SELECT COUNT(*) FROM medicao_documentos WHERE medicao_id=? AND tipo_doc='fatura'",
                        (mid,)).fetchone()[0] == 0


def test_sem_espelho_a_assinada_feita_a_mao_e_recusada_e_o_envio_bloqueia(conn, usuario, consorcio14, _texto):  # noqa: F811
    """Rodada 3 da PR #179 (S1, decisao de produto de 15/10/2026): medicao de consorcio sem espelho vigente nao gera FRC
    e a feita a' mao nao vale (a marca "sem formulario" nao existe no cadastro): o anexo e o envio ficam bloqueados."""
    cid, _ = consorcio14
    mid = medicoes.criar(conn, 1, {"contrato_id": cid, "numero": 50, "tipo": "execucao", "competencia": "2026-09"})
    assert frc.situacao(conn, mid)["bloqueio"] == "A medição não tem espelho vigente com as notas esperadas."
    with pytest.raises(ValueError, match="não tem espelho vigente com as notas esperadas: sem ele a dashboard não gera"):
        _assinar(conn, usuario, mid, _texto, "FRC 9 feita a mão.pdf", texto="FRC 00009 qualquer")
    assert conn.execute("SELECT COUNT(*) FROM medicao_documentos WHERE medicao_id=? AND tipo_doc='fatura'",
                        (mid,)).fetchone()[0] == 0
    a = frc.assinada(conn, mid)
    assert (a["controla"], a["documento"], a["recusa"]) == (True, None, frc.SEM_ESPELHO)
    item, faltas = _item_frc(conn, cid, mid)
    assert (item["estado"], item["bloqueia"], item["detalhe"]) == ("falta", True, frc.SEM_ESPELHO)
    assert not [f for f in faltas if "FRC" in f]           # nao e' pendencia justificavel: e' bloqueio
    c = protocolo._contrato(conn, cid)
    o = dict(conn.execute("SELECT * FROM orgaos WHERE id=?", (c["orgao_id"],)).fetchone())
    m = dict(conn.execute("SELECT * FROM medicoes WHERE id=?", (mid,)).fetchone())
    env = notas_envio._envio(conn, c, "sinfra", o, [m], None, __import__("datetime").date(2026, 10, 9))
    assert env["situacao"] == "bloqueado" and "Sem espelho para a FRC" in env["aviso"]
    # anexada antes desta regra (sem vinculo) tambem nao vale nem vai no e-mail
    conn.execute("INSERT INTO medicao_documentos (medicao_id, fonte, tipo_doc, nome, item_id, arquivo) VALUES "
                 "(?, 'operacao', 'fatura', 'velha.pdf', 'velha', 'notas/velha.pdf')", (mid,))
    conn.commit()
    assert frc.assinada(conn, mid)["documento"] is None
    assert not [d for d in protocolo._docs_nota(conn, [mid]) if d["tipo_doc"] == "fatura"]
    m = dict(conn.execute("SELECT * FROM medicoes WHERE id=?", (mid,)).fetchone())
    p = {"id": protocolo.criar(conn, 1, cid, "nota", [mid], None)}
    p = dict(conn.execute("SELECT * FROM protocolos WHERE id=?", (p["id"],)).fetchone())
    assert protocolo.frc_recusadas(conn, p) == [f"{medicoes.titulo(m)}: {frc.SEM_ESPELHO}"]


def test_baixar_frc_desatualizada_e_recusado(client_admin, _texto):  # noqa: F811
    """Issue #169: o download (botao da etapa 3 e link da etapa 4) nao entrega a FRC que ja' nao bate com as notas."""
    from app import db
    c = db.get_conn()
    try:
        cadastro_aplicar.aplicar(c, 1, cadastro_arquivos.preparar(CADASTRO_EXEMPLO), ensaio=False)
        _semear_frc(c)
        u = {"id": 1, "email": "admin@lotufo.com"}
        _, mid = _medicao14(c, u)
        _anexar(c, u, mid, _texto, NOTAS_FRC)
        for formato in ("pdf", "xlsx"):
            assert client_admin.get(f"/api/medicoes/{mid}/frc/{formato}").status_code == 200
        c.execute("UPDATE frcs SET notas='[]' WHERE medicao_id=? AND substituida_em IS NULL", (mid,))
        c.commit()
        for formato in ("pdf", "xlsx"):
            r = client_admin.get(f"/api/medicoes/{mid}/frc/{formato}")
            assert r.status_code == 409 and r.json()["detail"] == (
                "As notas mudaram depois desta FRC: gere de novo na etapa 3 (Nota fiscal) antes de baixar")
    finally:
        c.close()


def test_consorcio_sem_formulario_e_obra_propria(conn, usuario):  # noqa: F811
    cad = cadastro_arquivos.preparar(CADASTRO_EXEMPLO)
    cadastro_aplicar.aplicar(conn, 1, {**cad, "frc": []}, ensaio=False)          # consorcio sem FRC no cadastro
    consorcio = conn.execute("SELECT id FROM contratos WHERE nome_curto='Obra Consórcio'").fetchone()[0]
    mid = medicoes.criar(conn, 1, {"contrato_id": consorcio, "numero": 50, "tipo": "execucao", "competencia": "2026-09"})
    assert frc.situacao(conn, mid) == {"disponivel": False, "consorcio": "Consórcio Exemplo", "motivo": frc.INDISPONIVEL}
    with pytest.raises(ValueError, match="ainda não disponível"):
        frc.gerar(conn, usuario, mid, 1)
    # rodada 3 da PR #179 (S1): consorcio sem a FRC no cadastro nao aceita a feita a' mao e o envio bloqueia
    with pytest.raises(ValueError, match="A FRC deste consórcio não está no cadastro"):
        notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto("assinada"), "FRC 9.pdf", tipo_doc="fatura")
    assert frc.assinada(conn, mid) == {"controla": True, "documento": None, "recusa": frc.SEM_CADASTRO, "frc": None}
    item, faltas = _item_frc(conn, consorcio, mid)
    assert (item["estado"], item["bloqueia"]) == ("falta", True) and not [f for f in faltas if "FRC" in f]
    c = protocolo._contrato(conn, consorcio)
    o = dict(conn.execute("SELECT * FROM orgaos WHERE id=?", (c["orgao_id"],)).fetchone())
    m = dict(conn.execute("SELECT * FROM medicoes WHERE id=?", (mid,)).fetchone())
    env = notas_envio._envio(conn, c, "sinfra", o, [m], None, __import__("datetime").date(2026, 10, 9))
    assert env["situacao"] == "bloqueado" and "FRC fora do cadastro" in env["aviso"]
    propria = conn.execute("SELECT id FROM contratos WHERE nome_curto='Obra Própria'").fetchone()[0]
    mid = medicoes.criar(conn, 1, {"contrato_id": propria, "numero": 50, "tipo": "execucao", "competencia": "2026-09"})
    assert frc.situacao(conn, mid)["disponivel"] is False and frc.situacao(conn, mid)["motivo"] is None
    # obra propria: a `fatura` entra como antes, sem vinculo e sem ler o PDF
    doc = notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto("fatura"), "Fatura.pdf", tipo_doc="fatura")["documento_id"]
    assert frc.assinada(conn, mid) == {"controla": False, "documento": {"id": doc, "nome": "Fatura.pdf"}, "recusa": None,
                                       "frc": None}


def test_rotas_da_frc(client_admin):
    """Obra propria: a FRC nao se aplica, baixar e' 404 e gerar e' 422 com o motivo."""
    from tests.conftest import novo_contrato
    cid = novo_contrato({"orgao": "SINFRA", "numero_instrumento": "IC 1/2026", "objeto": "Obra",
                         "valor_original": "100.000,00"})
    mid = client_admin.post(f"/api/contratos/{cid}/medicoes",
                            json={"numero": 1, "tipo": "execucao", "competencia": "2026-08"}).json()["id"]
    r = client_admin.get(f"/api/medicoes/{mid}/frc")
    assert r.status_code == 200 and r.json()["disponivel"] is False
    assert client_admin.get(f"/api/medicoes/{mid}/frc/pdf").status_code == 404
    assert client_admin.get(f"/api/medicoes/{mid}/frc/doc").status_code == 404
    r = client_admin.post(f"/api/medicoes/{mid}/frc", json={"numero": 22})
    assert r.status_code == 422 and "consórcio" in r.json()["detail"]
    assert client_admin.post(f"/api/medicoes/{mid}/frc", json={"numero": 0}).status_code == 422
    assert client_admin.get("/api/medicoes/99999/frc").status_code == 404


def test_numero_digitado_abaixo_da_semente_reproduz_a_frc_ja_emitida(conn, usuario, _texto):  # noqa: F811
    """Consorcio cuja ultima FRC antes do sistema e' alta (51) e cujo formulario tem outro rotulo e rodape: a
    gerada sozinha sai 52, e o numero digitado 50 (uma FRC emitida fora de ordem) a regera com o texto certo."""
    cadastro_aplicar.aplicar(conn, 1, cadastro_arquivos.preparar(CADASTRO_EXEMPLO), ensaio=False)
    _semear_frc(conn, ultimo=51, rotulo_empresa="EMPRESA CONSORCIADA",
                descricao={"execucao": "{n}° MEDIÇÃO DA OBRA DE EXEMPLO, PERÍODO {inicio} A {fim}. DADOS BANCÁRIOS: "
                                       "BANCO EXEMPLO AGENCIA: 0001-0 CONTA CORRENTE: 00000-0"})
    cid = conn.execute("SELECT id FROM contratos WHERE nome_curto='Obra Consórcio'").fetchone()[0]
    eid = espelhos.gerar(conn, usuario, cid, 28, "2026-08", {"medicao": _soma(3)})
    mid = conn.execute("SELECT m.id FROM espelho_medicoes em JOIN medicoes m ON m.id=em.medicao_id "
                       "WHERE em.espelho_id=? AND m.tipo='execucao'", (eid,)).fetchone()[0]
    medicoes.marcar_espelho_enviado(conn, 1, mid)
    r = _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    assert r[-1]["frc_gerada"] == 52
    frc.gerar(conn, usuario, mid, 50)
    f = conn.execute("SELECT * FROM frcs WHERE medicao_id=? AND substituida_em IS NULL", (mid,)).fetchone()
    assert (f["numero"], f["bruto_centavos"], f["liquido_centavos"]) == (50, 1000000007, 924000007)
    t = " ".join(pypdf.PdfReader(io.BytesIO(frc.arquivo(conn, mid, "pdf").read_bytes())).pages[0].extract_text().split())
    for trecho in ("00050", "CONSÓRCIO EXEMPLO", "EMPRESA CONSORCIADA", "ENGENHARIA BETA S.A.",
                   "28° MEDIÇÃO DA OBRA DE EXEMPLO, PERÍODO 01/08/2026 A 31/08/2026", "AGENCIA: 0001-0", "9.240.000,07"):
        assert trecho in t, trecho
